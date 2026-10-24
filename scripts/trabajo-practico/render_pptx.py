"""Genera la exposición del 17 oct 2026: fundamentos del autocorrelador de intensidad y
propuesta de un diseño de comparación simple (técnica frente al dato del fabricante).

Uso: python render_pptx.py <salida.pptx> <carpeta_figuras>
El contenido sigue trabajo-practico/estructura-exposicion-no-estadistica.md.
"""
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ------------------------------------------------------------------ estilo
# Bahnschrift: tipografía técnica tipo DIN que trae Windows 10/11 (no hay que instalarla).
HEAD = "Bahnschrift SemiBold"
BODY = "Bahnschrift"
LIGHT = "Bahnschrift Light"
MATH = "Cambria Math"

DARK = RGBColor(0x0B, 0x10, 0x20)     # fondo de portada y cierre
INK = RGBColor(0x0E, 0x17, 0x26)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
RED = RGBColor(0xD7, 0x26, 0x3D)      # fundamental, ≈800 nm
VIOLET = RGBColor(0x5B, 0x3F, 0xD9)   # segundo armónico, ≈400 nm
VIOLET_L = RGBColor(0xA9, 0x9B, 0xFF)
RED_T = RGBColor(0xFD, 0xEC, 0xEE)
VIO_T = RGBColor(0xEF, 0xEC, 0xFF)
GRAY_T = RGBColor(0xF2, 0xF4, 0xF7)
LINE = RGBColor(0xD5, 0xDA, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = 13.333, 7.5
MX = 0.6
TOP = 1.75
FOOT = "Autocorrelador de intensidad · Diseño de Experimentos Avanzados · UNAL 2026-II"

TOKEN = re.compile(r"\*\*|\^\{|_\{|\}")


def parse(text):
    """Marcado mínimo: **negrita**, ^{superíndice}, _{subíndice}."""
    out, bold, mode, pos = [], False, "", 0
    for m in TOKEN.finditer(text):
        tok = m.group()
        if tok == "}" and not mode:
            continue
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, mode))
        pos = m.end()
        if tok == "**":
            bold = not bold
        elif tok == "^{":
            mode = "sup"
        elif tok == "_{":
            mode = "sub"
        else:
            mode = ""
    if pos < len(text):
        out.append((text[pos:], bold, mode))
    return out


def runs(par, text, size, color=INK, font=BODY, bold=False, italic=False):
    for t, b, mode in parse(text):
        r = par.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.name = font
        r.font.bold = False
        if bold or b:
            r.font.name = HEAD if font in (BODY, LIGHT) else font
            r.font.bold = font not in (BODY, LIGHT)
        r.font.italic = italic
        r.font.color.rgb = color
        if mode:
            r.font._element.set("baseline", "30000" if mode == "sup" else "-25000")


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, margin=0.0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tf


def text(slide, x, y, w, h, lines, size=17, color=INK, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=6, bold=False, italic=False, line_spacing=None):
    """lines: str o lista de str; cada elemento es un párrafo."""
    tf = box(slide, x, y, w, h, anchor)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line_spacing:
            p.line_spacing = line_spacing
        runs(p, ln, size, color, font, bold, italic)
    return tf


def bullets(slide, x, y, w, h, items, size=17, color=INK, space=9, marker=RED):
    """Viñetas con un pequeño cuadrado de color (el marcador es un carácter de viñeta real)."""
    tf = box(slide, x, y, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.28))))
        pPr.set("indent", str(-int(Inches(0.28))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(marker)})
        clr.append(srgb)
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        runs(p, it, size, color)
    return tf


def rect(slide, x, y, w, h, fill=GRAY_T, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.06, dash=False, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, title, body, fill=GRAY_T, tcolor=VIOLET, size=15, tsize=17,
         body_color=INK):
    rect(slide, x, y, w, h, fill)
    text(slide, x + 0.25, y + 0.2, w - 0.5, 0.45, title, tsize, tcolor, HEAD)
    if isinstance(body, list):
        bullets(slide, x + 0.25, y + 0.72, w - 0.5, h - 0.85, body, size, body_color, space=5,
                marker=tcolor)
    else:
        text(slide, x + 0.25, y + 0.72, w - 0.5, h - 0.85, body, size, body_color)


def equation(slide, x, y, w, h, lines, size=22, fill=VIO_T, caption=None):
    rect(slide, x, y, w, h, fill)
    ch = 0.35 if caption else 0
    text(slide, x + 0.2, y + 0.1, w - 0.4, h - 0.2 - ch, lines, size, INK, MATH, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE, space=4)
    if caption:
        text(slide, x + 0.2, y + h - 0.42, w - 0.4, 0.32, caption, 11.5, MUTED,
             align=PP_ALIGN.CENTER)


def quote(slide, x, y, w, h, q, source, size=15):
    rect(slide, x, y, w, h, GRAY_T)
    text(slide, x + 0.3, y + 0.18, w - 0.6, h - 0.62, "“" + q + "”", size, INK, LIGHT,
         italic=True)
    text(slide, x + 0.3, y + h - 0.45, w - 0.6, 0.3, "— " + source, 12, MUTED)


def stat(slide, x, y, w, big, label, color=RED, size=40):
    text(slide, x, y, w, 0.8, big, size, color, HEAD)
    text(slide, x, y + 0.8, w, 0.7, label, 13.5, MUTED)


def pending(slide, x, y, w, h, msg):
    rect(slide, x, y, w, h, WHITE, RED, dash=True)
    text(slide, x + 0.2, y, w - 0.4, h, "Pendiente: " + msg, 13, RED, anchor=MSO_ANCHOR.MIDDLE)


def image(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def arrow(slide, x1, y1, x2, y2, color=INK, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2),
                                   Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def table(slide, x, y, w, rows, widths, size=14, head_fill=INK, row_h=0.42, pending_col=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # sin estilo
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = head_fill if i == 0 else (GRAY_T if i % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            is_pending = pending_col is not None and j == pending_col and i > 0
            color = WHITE if i == 0 else (RED if is_pending else INK)
            runs(p, val, size, color, HEAD if i == 0 else BODY, italic=is_pending)
    return tbl


# ------------------------------------------------------------------ marcos
class Deck:
    def __init__(self, figdir):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.fig = lambda n: os.path.join(figdir, n + ".png")
        self.n = 0

    def _bg(self, slide, color):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = color

    def dark(self, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        self._bg(s, DARK)
        self.n += 1
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def content(self, kicker, title, notes="", refs=None):
        s = self.prs.slides.add_slide(self.blank)
        self._bg(s, WHITE)
        self.n += 1
        text(s, MX, 0.36, 10.5, 0.35, kicker.upper(), 12.5, VIOLET, HEAD)
        text(s, MX, 0.66, 11.6, 0.95, title, 30, INK, HEAD, anchor=MSO_ANCHOR.TOP)
        image(s, self.fig("icono"), SW - MX - 0.95, 0.38, w=0.95)
        text(s, MX, SH - 0.42, 9.0, 0.3, FOOT, 10, MUTED)
        if refs:
            text(s, MX + 6.2, SH - 0.42, 5.6, 0.3, refs, 10, MUTED, align=PP_ALIGN.RIGHT)
        text(s, SW - MX - 0.6, SH - 0.42, 0.6, 0.3, str(self.n), 10, MUTED, align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s


def icono(path):
    """Pequeño pulso (rojo) con su segundo armónico (violeta): motivo de cada diapositiva."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    t = np.linspace(-60, 60, 3000)
    fig, ax = plt.subplots(figsize=(1.9, 0.75))
    fig.patch.set_alpha(0)
    env = 1 / np.cosh(t / 12)
    ax.plot(t, env * np.cos(2.35 * t), color="#D7263D", lw=1.3)
    ax.plot(t, 0.55 * env ** 2 * np.cos(4.7 * t) - 1.45, color="#5B3FD9", lw=1.1)
    ax.axis("off")
    fig.savefig(path, transparent=True, dpi=220, bbox_inches="tight", pad_inches=0.01)
    plt.close(fig)


# ------------------------------------------------------------------ diapositivas
def build(out, figdir):
    icono(os.path.join(figdir, "icono.png"))
    d = Deck(figdir)
    F = d.fig

    # 1 · Portada ------------------------------------------------------------
    s = d.dark("Presentarse y presentar la propuesta en una frase: vamos a calibrar un "
               "autocorrelador de intensidad comparando el ancho de pulso que mide con el valor "
               "que reporta el fabricante del láser. Todo lo que sigue es la física que hace "
               "falta para entender por qué esa comparación es el primer paso correcto.")
    image(s, F("portada"), 6.9, 0.9, w=6.2)
    text(s, MX, 1.0, 7.0, 0.4, "DISEÑO DE EXPERIMENTOS AVANZADOS · TRABAJO PRÁCTICO", 13, VIOLET_L,
         HEAD)
    text(s, MX, 1.55, 6.9, 2.0, "Calibración de un autocorrelador de intensidad", 44, WHITE, HEAD,
         line_spacing=0.95)
    text(s, MX, 3.75, 6.4, 1.2, "Fundamentos físicos de la medición de pulsos ultracortos y "
         "propuesta de un diseño de comparación simple", 19, RGBColor(0xC9, 0xD1, 0xDE), LIGHT)
    text(s, MX, 5.35, 7.0, 0.9, ["Juan Esteban Rodríguez Ochoa · Sebastián Zapata Henao",
                                 "Profesora: Faviana Gutiérrez Rôa"], 15, WHITE, BODY, space=4)
    text(s, MX, 6.4, 8.0, 0.4, "Universidad Nacional de Colombia · Sede Medellín · 17 de octubre "
         "de 2026", 13, RGBColor(0x9A, 0xA4, 0xB5), BODY)

    # 2 · Propuesta y agenda ----------------------------------------------------
    s = d.content("Qué proponemos", "¿Qué vamos a proponer y cómo llegamos ahí?",
                  "Anunciar desde el principio adónde vamos: así cada bloque de física se escucha "
                  "como parte del argumento. La propuesta es un experimento de comparación simple: "
                  "patrón (dato del fabricante) frente a técnica (autocorrelador de intensidad). Los "
                  "cinco bloques siguientes construyen las razones; el sexto presenta el diseño.")
    rect(s, MX, TOP, 5.4, 4.9, VIO_T)
    text(s, MX + 0.35, TOP + 0.3, 4.7, 0.4, "PROPUESTA", 13, VIOLET, HEAD)
    text(s, MX + 0.35, TOP + 0.8, 4.75, 2.6, "Comparar el ancho de pulso medido con el "
         "autocorrelador de intensidad contra el valor del fabricante del láser.", 24, INK, HEAD,
         line_spacing=1.05)
    text(s, MX + 0.35, TOP + 3.25, 4.75, 1.5, ["**Para qué:** calibrar la técnica y "
                                               "diagnosticar el estado del láser.",
                                               "**Diseño:** comparación simple, patrón frente "
                                               "a técnica."], 16, INK, space=8)
    agenda = ["¿Qué queremos medir?", "¿De dónde sale el pulso y qué dice el fabricante?",
              "¿Cómo funciona el autocorrelador?", "¿Cómo se pasa de la traza al ancho del pulso?",
              "¿Qué puede cambiar el ancho medido?", "¿Por qué empezar por calibrar?"]
    for i, a in enumerate(agenda):
        y = TOP + 0.05 + i * 0.81
        circ = rect(s, 6.55, y, 0.6, 0.6, RED if i == 5 else INK, shape=MSO_SHAPE.OVAL)
        text(s, 6.55, y, 0.6, 0.6, str(i + 1), 18, WHITE, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, 7.4, y, 5.3, 0.6, a, 18, RED if i == 5 else INK, HEAD if i == 5 else BODY,
             anchor=MSO_ANCHOR.MIDDLE)
        del circ

    # 3 · Escala de tiempos ---------------------------------------------------
    s = d.content("1 · ¿Qué queremos medir?",
                  "¿Por qué no basta un detector para medir el pulso?",
                  "1 fs es 10^-15 s. Un fotodiodo rápido responde en decenas de picosegundos y un "
                  "osciloscopio rápido en picosegundos: el pulso es de 100 a 1000 veces más corto. "
                  "El detector solo ve la energía total del pulso, no su forma. Por eso la idea "
                  "central de la técnica es que el pulso se mide a sí mismo. Dejar sembrada la "
                  "consecuencia: un instrumento así no se verifica con un cronómetro, necesita una "
                  "referencia externa; esa referencia es el fabricante (bloque 6). Los rangos de los "
                  "detectores son órdenes de magnitud típicos, no especificaciones de un equipo.",
                  "Escalas: órdenes de magnitud típicos")
    image(s, F("escala"), 0.95, TOP - 0.05, w=11.4)
    cards = [("Lo que ve un detector", "Solo la energía integrada del pulso, nunca su forma.",
              GRAY_T, VIOLET),
             ("La idea de la técnica", "El pulso tiene que medirse **a sí mismo**.", GRAY_T, VIOLET),
             ("La consecuencia", "Un instrumento así no se verifica con un cronómetro: necesita "
              "una **referencia externa**.", RED_T, RED)]
    for i, (t_, b_, f_, c_) in enumerate(cards):
        card(s, MX + i * 4.1, 5.2, 3.85, 1.6, t_, b_, f_, c_, size=16)

    # 4 · Qué es un pulso ---------------------------------------------------
    s = d.content("1 · ¿Qué queremos medir?", "¿Qué es exactamente un pulso ultracorto?",
                  "El campo es una portadora óptica (a 800 nm oscila cada 2,7 fs) modulada por una "
                  "envolvente. Lo que se mide es la intensidad, el cuadrado de la envolvente, y el "
                  "ancho de pulso se define como su ancho a media altura (FWHM). Las dos formas de "
                  "referencia en la literatura son la gaussiana y la secante hiperbólica al cuadrado; "
                  "la forma importa porque, como veremos, el ancho que reporta el autocorrelador "
                  "depende de cuál se suponga.", "Diels y Rudolph (2006), cap. 1")
    equation(s, MX, TOP, 5.6, 1.55, ["E(t) = Re{ A(t) · e^{iω₀t} }", "I(t) ∝ |A(t)|²"], 24)
    bullets(s, MX, TOP + 1.85, 5.6, 3.2, [
        "**A(t):** envolvente lenta; **ω₀:** portadora óptica (a 800 nm, un ciclo dura 2,7 fs).",
        "**Ancho de pulso τ_{p}:** ancho a media altura (FWHM) de la intensidad I(t).",
        "**Formas de referencia:** gaussiana y secante hiperbólica al cuadrado (sech²).",
    ], 17)
    image(s, F("pulso"), 6.55, TOP - 0.05, w=6.2)

    # 5 · Fourier -------------------------------------------------------------
    s = d.content("1 · ¿Qué queremos medir?", "¿Por qué un pulso corto necesita un espectro ancho?",
                  "Tiempo y frecuencia son pares de Fourier: un pulso cinco veces más corto tiene un "
                  "espectro cinco veces más ancho (figura). El producto Δν·τp tiene un mínimo que "
                  "depende de la forma: 0,441 para la gaussiana y 0,315 para la sech² (Diels y "
                  "Rudolph, tabla 1.1, p. 10). Un pulso que alcanza ese mínimo se llama limitado por "
                  "transformada; si tiene chirp, su espectro es más ancho que el límite (p. 18). "
                  "Ejemplo: 100 fs gaussianos a 800 nm necesitan Δν = 4,41 THz, unos 9,4 nm. Por eso "
                  "el espectro es un segundo dato para comparar con el fabricante.",
                  "Diels y Rudolph (2006), tabla 1.1, p. 10 y p. 18")
    image(s, F("fourier"), MX, TOP - 0.05, w=6.6)
    equation(s, 7.55, TOP, 5.2, 1.0, "Δν · τ_{p} ≥ K", 26)
    table(s, 7.55, TOP + 1.2, 5.2, [["Forma del pulso", "K (mínimo)"], ["Gaussiana", "0,441"],
                                   ["sech²", "0,315"]], [3.2, 2.0], 15, row_h=0.42)
    stat(s, 7.55, TOP + 2.75, 5.2, "100 fs → Δλ ≈ 9,4 nm", "pulso gaussiano limitado por "
         "transformada, centrado en 800 nm", RED, 30)
    text(s, 7.55, TOP + 4.25, 5.2, 0.8, "El espectro es un **segundo dato** para comparar con el "
         "fabricante y para validar el ancho medido.", 16)

    # 6 · Mode-locking -------------------------------------------------------
    s = d.content("2 · ¿De dónde sale el pulso?",
                  "¿Cómo genera pulsos un láser de modo amarrado?",
                  "La cavidad admite modos longitudinales separados c/(2nL) (Diels y Rudolph, ec. "
                  "5.1, p. 278). Si muchos modos oscilan con la misma fase, se suman en un tren de "
                  "pulsos que se repite cada tiempo de ida y vuelta; con fases al azar solo hay "
                  "ruido (figura, 15 modos simulados). Más modos en fase, es decir más ancho de "
                  "banda, dan un pulso más corto. La forma que se supone al medir suele ser sech² o "
                  "gaussiana (p. 459); la sech² se asocia a la teoría de Haus (1975) y a los "
                  "solitones (Diels y Rudolph, pp. 296 y 386).",
                  "Diels y Rudolph (2006), ec. 5.1, p. 278 y p. 459")
    equation(s, MX, TOP, 5.6, 1.5, ["Δν = c / (2nL)", "T = 2nL / c"], 24,
             caption="separación entre modos y periodo del tren")
    bullets(s, MX, TOP + 1.8, 5.6, 3.2, [
        "Muchos modos longitudinales oscilan con **fase fija**: su suma es un tren de pulsos.",
        "**Más modos en fase** (más ancho de banda) → pulso más corto.",
        "Al medir se supone una forma: \"generally a sech² or a Gaussian shape\" (Diels y "
        "Rudolph, p. 459).",
    ], 17)
    image(s, F("modos"), 6.55, TOP - 0.05, w=6.2)

    # 7 · Fabricante -------------------------------------------------------------
    # Fuente: Calmar Laser, hoja de datos Carmel X-780, Rev. E (10/1/2025), págs. 2 y 3.
    # Variante asumida: X-780. Las demás (X-1550, X-780/1550, X-1030, X-515) están en
    # notas/laser-carmel-x-ficha-tecnica.md; el modelo exacto se lee en la etiqueta del equipo.
    s = d.content("2 · ¿Qué dice el fabricante?", "¿Qué especifica el fabricante y cómo lo midió?",
                  "El láser es un Calmar Laser Carmel X, de fibra, con mode-locking pasivo por "
                  "absorbedor saturable. Su hoja de datos (variante X-780, Rev. E, 2025) define el "
                  "patrón: el ancho de pulso es un máximo garantizado, < 90 fs, no un valor "
                  "nominal, y el fabricante lo obtuvo con la misma técnica nuestra: una traza de "
                  "autocorrelación de segundo armónico, deconvolucionada suponiendo forma sech² "
                  "con factor 0,65 (nota 2 de la tabla de especificaciones). Por eso la prueba es "
                  "de una cola y hay que usar la misma forma al deconvolucionar. El espectro "
                  "típico (λc = 781,3 nm, Δλ = 10,4 nm RMS) y la traza típica (FWHM ≈ 0,14 ps, es "
                  "decir ≈ 90 fs) son del modelo, no de este equipo. Falta leer el modelo exacto "
                  "en la etiqueta del controlador (CFL-04, -10 o -15RFF; o la variante de 1550 nm) "
                  "y buscar si vino con hoja de pruebas de fábrica, que sería el patrón más fuerte.",
                  "Calmar Laser, hoja de datos Carmel X-780, Rev. E (2025), págs. 2–3")
    table(s, MX, TOP, 7.0, [
        ["Parámetro", "Carmel X-780 (hoja de datos)", "Para la comparación"],
        ["Ancho de pulso", "**< 90 fs** (máximo)", "Prueba de una cola: H_{A}: μ_{τ} > 90 fs"],
        ["Técnica del fabricante", "Autocorrelación de 2.º armónico", "La misma que la nuestra"],
        ["Forma supuesta", "sech², factor 0,65", "τ_{p} = 0,65 · τ_{AC}: usar la misma"],
        ["Longitud de onda central", "780 ± 3 nm", "Espectro típico: λ_{c} = 781,3 nm"],
        ["Ancho espectral (FWHM)", "8 – 10 nm", "9 nm y 90 fs → Δν·τ_{p} ≈ 0,40"],
        ["Tasa de repetición", "50 u 80 MHz, según modelo", "Leer en la etiqueta del equipo"],
        ["Potencia media", "0,25 / 1,0 / 1,5 W (-04/-10/-15)", "Leer en la etiqueta del equipo"],
    ], [2.25, 2.45, 2.3], 13, row_h=0.5)
    rect(s, 7.85, TOP, 4.88, 1.3, RED_T)
    text(s, 8.1, TOP + 0.15, 4.45, 0.35, "EL PATRÓN DEL EXPERIMENTO", 12.5, RED, HEAD)
    text(s, 8.1, TOP + 0.5, 4.45, 0.8, "Un máximo garantizado, medido con autocorrelación y "
         "forma sech².", 16, INK, HEAD)
    image(s, F("fabricante_traza"), 8.45, TOP + 1.45, w=3.65)   # 586x421 px → 2.62 in de alto
    text(s, 7.85, TOP + 4.12, 4.88, 0.55, "Traza de autocorrelación típica del X-780 (hoja de "
         "datos, pág. 3). FWHM ≈ 0,14 ps; con 0,65 da ≈ 90 fs.", 11.5, MUTED)

    # 8 · Dispersión ------------------------------------------------------------
    s = d.content("2 · ¿De dónde sale el pulso?",
                  "¿Qué le pasa al pulso antes de llegar al autocorrelador?",
                  "El vidrio retrasa unas frecuencias más que otras: el pulso sale con chirp y más "
                  "largo, aunque el espectro sea el mismo. Para un pulso gaussiano sin chirp, el "
                  "ancho de salida sigue la fórmula de la diapositiva: es la ec. 1.125 de Diels y "
                  "Rudolph (p. 34) escrita con el FWHM. El ejemplo usa 1 cm de BK7 a 800 nm: k'' ≈ "
                  "45 fs²/mm, calculado con n'' = 0,05 µm⁻² de la tabla 2.1 (p. 62). Un pulso de 100 "
                  "fs casi no cambia, uno de 30 fs crece un 70 %. Mensaje para la comparación: hay "
                  "que contar todo el vidrio entre el láser y el cristal, incluido el divisor de haz "
                  "y la lente del autocorrelador.",
                  "Diels y Rudolph (2006), ec. 1.125, p. 34; tabla 2.1, p. 62")
    image(s, F("gdd"), MX, TOP - 0.05, w=6.3)
    equation(s, 7.3, TOP, 5.43, 1.35, "τ_{sal} = τ_{ent} · √(1 + (4 ln2 · GDD / τ_{ent}²)²)", 20,
             caption="pulso gaussiano sin chirp; GDD = dispersión de retardo de grupo")
    text(s, 7.3, TOP + 1.6, 5.4, 0.4, "1 cm de vidrio BK7 a 800 nm (GDD ≈ 450 fs²):", 15, MUTED)
    stat(s, 7.3, TOP + 1.95, 2.6, "100 → 101 fs", "casi no cambia", INK, 28)
    stat(s, 10.1, TOP + 1.95, 2.6, "30 → 51 fs", "crece un 70 %", RED, 28)
    rect(s, 7.3, TOP + 3.55, 5.43, 1.45, VIO_T)
    text(s, 7.5, TOP + 3.65, 5.0, 1.3, "Todo el vidrio entre el láser y el cristal cuenta "
         "(divisor de haz, lentes, ventanas): sin descontarlo se mide un pulso **más largo** que el "
         "que entrega el láser.", 16, anchor=MSO_ANCHOR.MIDDLE)

    # 9 · Principio --------------------------------------------------------------
    s = d.content("3 · ¿Cómo funciona el autocorrelador?",
                  "¿Cómo se mide algo más rápido que cualquier detector?",
                  "Tres pasos: dividir el pulso en dos copias, retrasar una de ellas, y combinarlas "
                  "en un medio no lineal que solo responde cuando las dos copias llegan juntas. El "
                  "detector es lento y solo integra esa señal; barriendo el retardo se recorre el "
                  "pulso con el pulso mismo. El retardo sale del camino recorrido: el factor 2 es "
                  "por la ida y vuelta en el retrorreflector, así que 1 µm de desplazamiento son "
                  "6,67 fs. La técnica es la que propuso Armstrong en 1967.",
                  "Armstrong (1967); Diels y Rudolph (2006), cap. 9")
    steps = [("1", "Dividir", "Un divisor de haz crea dos copias idénticas del pulso."),
             ("2", "Retrasar", "Una copia recorre un camino más largo y llega un tiempo τ después."),
             ("3", "Combinar", "Un medio no lineal da señal solo cuando las copias se superponen; "
              "el detector lento la integra.")]
    for i, (n_, t_, b_) in enumerate(steps):
        x = MX + i * 4.1
        rect(s, x, TOP, 3.85, 2.75, GRAY_T)
        rect(s, x + 0.25, TOP + 0.25, 0.65, 0.65, RED if i == 2 else INK, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.25, TOP + 0.25, 0.65, 0.65, n_, 20, WHITE, HEAD, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE)
        text(s, x + 1.05, TOP + 0.3, 2.6, 0.55, t_, 21, INK, HEAD, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, TOP + 1.1, 3.35, 1.6, b_, 16)
        if i < 2:
            arrow(s, x + 3.88, TOP + 1.37, x + 4.07, TOP + 1.37, MUTED, 2)
    equation(s, MX, TOP + 3.1, 5.0, 1.6, "τ = 2Δx / c", 30, caption="Δx: desplazamiento del "
             "retrorreflector")
    stat(s, 6.0, TOP + 3.15, 3.2, "6,67 fs", "de retardo por cada 1 µm que se mueve el "
         "retrorreflector", RED, 40)
    text(s, 9.5, TOP + 3.2, 3.25, 1.6, "Barriendo τ se recorre el pulso **con el pulso mismo**.",
         19, INK, anchor=MSO_ANCHOR.MIDDLE)

    # 10 · Montaje ----------------------------------------------------------------
    s = d.content("3 · ¿Cómo funciona el autocorrelador?", "¿Cómo está armado el autocorrelador?",
                  "Esquema genérico de un autocorrelador de barrido tipo Michelson. El rojo es la luz "
                  "del láser (fundamental) y el violeta el segundo armónico que se genera en el "
                  "cristal. Se reemplaza por el esquema del montaje real cuando tengamos sus "
                  "componentes. En geometría colineal los dos haces viajan superpuestos y aparecen "
                  "franjas y fondo; en la no colineal se cruzan en ángulo y la señal de suma sale "
                  "sola, sin fondo.", "Donnelly y Grossman (1998); Sala et al. (1980)")
    y0 = TOP + 0.35
    blocks = [(MX, "Láser", INK), (2.75, "Divisor\nde haz", INK)]
    for x, lab, col in blocks:
        rect(s, x, y0 + 0.75, 1.65, 1.0, col)
        text(s, x, y0 + 0.75, 1.65, 1.0, lab, 15, WHITE, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    arrow(s, MX + 1.67, y0 + 1.25, 2.73, y0 + 1.25, RED, 3)
    for k, (lab, yy) in enumerate([("Brazo fijo", y0), ("Brazo móvil\nretrorreflector · Δx", y0 + 1.5)]):
        rect(s, 4.95, yy, 2.4, 1.0, GRAY_T, LINE)
        text(s, 4.95, yy, 2.4, 1.0, lab, 14, INK, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        arrow(s, 4.42, y0 + 1.25, 4.93, yy + 0.5, RED, 3)
        arrow(s, 7.37, yy + 0.5, 7.88, y0 + 1.25, RED, 3)
    chain = [(7.9, "Lente", INK), (9.0, "Cristal\nSHG", VIOLET), (10.1, "Filtro", INK),
             (11.2, "Detector", INK)]
    for i, (x, lab, col) in enumerate(chain):
        rect(s, x, y0 + 0.75, 0.95, 1.0, col)
        text(s, x - 0.05, y0 + 0.75, 1.05, 1.0, lab, 13, WHITE, HEAD, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE)
        if i < 3:
            arrow(s, x + 0.97, y0 + 1.25, chain[i + 1][0] - 0.02, y0 + 1.25,
                  RED if i == 0 else VIOLET, 3)
    text(s, 11.0, y0 + 1.85, 1.75, 0.6, "señal(τ)", 14, VIOLET, HEAD, PP_ALIGN.CENTER)
    tf = box(s, MX, 4.78, 6.8, 0.35)
    par = tf.paragraphs[0]
    for t_, c_ in [("■ ", RED), ("Fundamental (≈ 800 nm)      ", MUTED), ("■ ", VIOLET),
                   ("Segundo armónico (≈ 400 nm)", MUTED)]:
        runs(par, t_, 13, c_)
    card(s, MX, 5.35, 5.9, 1.5, "Colineal", "Haces superpuestos: aparecen **franjas** y **fondo**; "
         "relación pico/fondo 3:1 u 8:1.", size=15)
    card(s, 6.83, 5.35, 5.9, 1.5, "No colineal", "Haces cruzados en ángulo: la señal de suma "
         "sale por el centro, **sin fondo**.", size=15)

    # 11 · Por qué no lineal --------------------------------------------------------
    s = d.content("3 · ¿Cómo funciona el autocorrelador?", "¿Por qué hace falta un proceso no lineal?",
                  "Con un detector lineal la señal es la autocorrelación del campo, y su "
                  "transformada de Fourier es el espectro (teorema de Wiener–Khinchin): no hay "
                  "información de duración que no dé ya un espectrómetro. Diels y Rudolph lo dicen "
                  "explícitamente en la p. 460. Para medir duración hay que multiplicar las dos "
                  "intensidades, y eso solo lo hace un proceso de segundo orden: generación de "
                  "segundo armónico en un cristal o absorción de dos fotones en un semiconductor.",
                  "Diels y Rudolph (2006), p. 460")
    for i, (ttl, eq, lines, fill, col) in enumerate([
        ("Detector lineal", "∫ |E(t) + E(t − τ)|² dt",
         ["Mide la autocorrelación del **campo**.", "Su transformada de Fourier es el espectro.",
          "**No da el ancho del pulso.**"], GRAY_T, MUTED),
        ("Proceso de segundo orden", "∫ I(t) · I(t − τ) dt",
         ["Mide la autocorrelación de **intensidad**.", "Solo hay señal si las copias se "
          "superponen.", "**Da el ancho del pulso.**"], VIO_T, VIOLET)]):
        x = MX + i * 6.23
        rect(s, x, TOP, 5.9, 3.35, fill)
        text(s, x + 0.3, TOP + 0.22, 5.3, 0.45, ttl, 20, col, HEAD)
        text(s, x + 0.3, TOP + 0.8, 5.3, 0.7, eq, 22, INK, MATH, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE)
        bullets(s, x + 0.3, TOP + 1.65, 5.3, 1.7, lines, 16, marker=col, space=5)
    quote(s, MX, TOP + 3.6, 12.13, 1.4, "A first-order field autocorrelation does not carry any "
          "other information than that provided by a spectrometer.", "Diels y Rudolph (2006), "
          "p. 460", 17)

    # 12 · SHG ------------------------------------------------------------------------
    s = d.content("3 · ¿Cómo funciona el autocorrelador?",
                  "¿Qué es la generación de segundo armónico (SHG)?",
                  "El término de segundo orden de la polarización oscila al doble de frecuencia: de "
                  "800 nm sale luz a 400 nm. La eficiencia depende del desajuste de fase Δk; en "
                  "cristales birrefringentes como BBO o KDP se anula girando el cristal (ajuste de "
                  "fase). El espesor es un compromiso: un cristal más grueso da más señal (crece "
                  "con L²) pero acepta menos ancho de banda, y si es demasiado largo la traza medida "
                  "sale más ancha que la real (Diels y Rudolph, p. 468). Desajuste de velocidad de "
                  "grupo a 800 nm: BBO 187 fs/mm, KDP 77 fs/mm (tabla 3.2, p. 176). Prueba práctica: "
                  "el espectro del armónico no debe salir más angosto que el fundamental al "
                  "cuadrado.", "Diels y Rudolph (2006), pp. 175–176 y 468")
    equation(s, MX, TOP, 6.0, 2.0, ["P = ε₀ (χ⁽¹⁾E + χ⁽²⁾E² + …)",
                                    "I_{2ω} ∝ L² · I_{ω}² · sinc²(ΔkL/2)"], 21,
             caption="Δk = k(2ω) − 2k(ω): ajuste de fase cuando Δk = 0")
    table(s, MX, TOP + 2.3, 6.0, [["Cristal (800 nm)", "Desajuste de velocidad de grupo"],
                                  ["BBO", "187 fs/mm"], ["KDP", "77 fs/mm"]], [2.6, 3.4], 15,
          row_h=0.42)
    text(s, MX, TOP + 3.65, 6.0, 0.3, "Diels y Rudolph, tabla 3.2, p. 176", 11.5, MUTED)
    rect(s, MX, TOP + 4.05, 6.0, 1.0, RED_T)
    text(s, MX + 0.2, TOP + 4.05, 5.6, 1.0, "Cristal demasiado largo → la traza sale **más "
         "ancha** que la real (p. 468).", 16, anchor=MSO_ANCHOR.MIDDLE)
    image(s, F("sinc"), 6.95, TOP - 0.05, w=5.8)

    # 13 · Trazas -------------------------------------------------------------------
    s = d.content("3 · ¿Cómo funciona el autocorrelador?", "¿Qué señal registra el detector?",
                  "Simulación para un pulso sech² de 20 fs. Sin fondo, la traza es la "
                  "autocorrelación de intensidad A(τ), siempre simétrica. En geometría colineal "
                  "promediada aparece un fondo y la relación pico/fondo es 3:1; si se resuelven las "
                  "franjas, la traza interferométrica tiene relación 8:1 (Diels y Rudolph, p. 462 y "
                  "fig. 9.1). Esas relaciones sirven para revisar la alineación antes de cada serie "
                  "de mediciones: si no se cumplen, algo está mal.",
                  "Diels y Rudolph (2006), p. 462 y fig. 9.1")
    image(s, F("trazas"), 1.15, TOP - 0.1, w=11.0)
    equation(s, MX, 5.55, 5.6, 1.3, "A(τ) = ∫ I(t) · I(t − τ) dt", 24,
             caption="autocorrelación de intensidad: siempre simétrica en τ")
    rect(s, 6.5, 5.55, 6.23, 1.3, VIO_T)
    text(s, 6.75, 5.55, 5.8, 1.3, "Las relaciones **3:1** y **8:1** son un control de "
         "alineación antes de cada serie de mediciones.", 17, anchor=MSO_ANCHOR.MIDDLE)

    # 14 · Deconvolución -----------------------------------------------------------
    s = d.content("4 · ¿De la traza al ancho del pulso?", "¿Cómo se obtiene τ_{p} a partir de la traza?",
                  "La traza no es el pulso: su ancho τAC es k veces el ancho del pulso, y k depende "
                  "de la forma que se suponga. Los valores son de la tabla II de Sala et al. (1980, "
                  "p. 993), que da τp/τAC para diez formas; Diels y Rudolph (tabla 9.1, p. 477) "
                  "confirman 1,414 y 1,543. La misma traza da 70,7 fs si se supone gaussiana y "
                  "64,8 fs si se supone sech²: un 9 % de diferencia. Para comparar con el fabricante "
                  "hay que usar la misma forma que él. Estos factores son de la autocorrelación de "
                  "intensidad; la interferométrica usa otros.",
                  "Sala et al. (1980), tabla II, p. 993; Diels y Rudolph (2006), tabla 9.1")
    equation(s, MX, TOP, 5.9, 1.0, "τ_{AC} = k · τ_{p}", 26)
    table(s, MX, TOP + 1.2, 5.9, [["Forma del pulso", "k = τ_{AC}/τ_{p}", "Δν·τ_{p}"],
                                  ["Cuadrada", "1", "0,886"], ["Gaussiana", "1,414", "0,441"],
                                  ["sech²", "1,543", "0,315"], ["Lorentziana", "2,000", "0,221"]],
          [2.3, 1.9, 1.7], 15, row_h=0.42)
    rect(s, MX, TOP + 3.55, 5.9, 1.45, RED_T)
    text(s, MX + 0.25, TOP + 3.55, 5.4, 1.45, "Para comparar con el fabricante hay que usar **la "
         "misma forma supuesta** que usó él.", 18, anchor=MSO_ANCHOR.MIDDLE)
    image(s, F("deconv"), 7.0, TOP - 0.05, w=5.6)
    text(s, 6.85, TOP + 4.15, 5.9, 0.9, "Misma traza (τ_{AC} = 100 fs), dos respuestas: "
         "**9 % de diferencia** solo por la forma supuesta.", 16)

    # 15 · Calibración del retardo -----------------------------------------------------
    s = d.content("4 · ¿De la traza al ancho del pulso?", "¿Cómo se calibra el eje de retardo?",
                  "El ancho se lee sobre el eje de retardo, así que un error de escala en ese eje "
                  "pasa completo al resultado. Tres formas de calibrarlo que aparecen en las "
                  "fuentes: mover el brazo una distancia conocida con una platina de precisión (Sala "
                  "et al., p. 992), insertar una lámina de vidrio de espesor conocido y medir el "
                  "corrimiento (Diels y Rudolph, fig. 9.6, p. 470), o contar franjas de la traza "
                  "interferométrica, cada una vale λ/c (Donnelly y Grossman, p. 683). También hay que "
                  "verificar que la escala sea lineal: Sala et al. reportan ±10 % en los dos tercios "
                  "centrales de su barrido (p. 995). Es la calibración interna, previa a la "
                  "comparación con el fabricante.",
                  "Sala et al. (1980); Diels y Rudolph (2006), fig. 9.6; Donnelly y Grossman (1998)")
    methods = [("Platina de precisión", "Mover el brazo una distancia conocida y medir cuánto se "
                "corre el pico.", "Sala et al., p. 992"),
               ("Lámina de vidrio", "Insertar en un brazo una lámina de espesor conocido y medir "
                "el corrimiento de la traza.", "Diels y Rudolph, fig. 9.6, p. 470"),
               ("Franjas de interferencia", "Contar franjas: cada una vale λ/c (2,67 fs a "
                "800 nm).", "Donnelly y Grossman, p. 683")]
    for i, (t_, b_, src) in enumerate(methods):
        x = MX + i * 4.1
        rect(s, x, TOP, 3.85, 2.9, GRAY_T)
        text(s, x + 0.25, TOP + 0.22, 3.4, 0.45, t_, 19, VIOLET, HEAD)
        text(s, x + 0.25, TOP + 0.8, 3.4, 1.5, b_, 16)
        text(s, x + 0.25, TOP + 2.4, 3.4, 0.35, src, 11.5, MUTED)
    stat(s, MX, TOP + 3.25, 4.2, "±10 %", "linealidad del barrido reportada por Sala et al. en "
         "los 2/3 centrales (p. 995)", VIOLET, 40)
    rect(s, 5.2, TOP + 3.25, 7.53, 1.75, RED_T)
    text(s, 5.45, TOP + 3.25, 7.1, 1.75, "Un error de escala en τ pasa **completo** al ancho "
         "medido: es la primera calibración, antes de comparar con el fabricante.", 18,
         anchor=MSO_ANCHOR.MIDDLE)

    # 16 · Límites ---------------------------------------------------------------------
    s = d.content("4 · ¿De la traza al ancho del pulso?",
                  "¿Qué no puede decirnos la autocorrelación de intensidad?",
                  "Tres límites con fuente. Uno: no da la fase; un pulso con chirp puede dar la "
                  "misma traza que uno más largo sin chirp (Monmayrant et al., p. 19). Dos: Chung y "
                  "Weiner construyen pulsos distintos con la misma autocorrelación y el mismo "
                  "espectro, cuyos anchos FWHM difieren del orden de 25 % (p. 665). Tres: si los "
                  "pulsos del tren no son iguales, el promedio produce un pico angosto sobre un "
                  "pedestal ancho; el pico no es el pulso, la duración real la indica el pedestal "
                  "(Ratner et al., p. 2874). La figura es una simulación de ese caso. Para este "
                  "proyecto basta con el ancho; FROG o SPIDER solo se mencionan.",
                  "Monmayrant et al. (2010); Chung y Weiner (2001); Ratner et al. (2012)")
    image(s, F("artefacto"), MX, TOP - 0.05, w=6.2)
    lim = [("Sin fase ni chirp", "\"A chirped pulse can give exactly the same autocorrelation "
            "signal as a longer Fourier-limited pulse\" (Monmayrant et al., p. 19)."),
           ("Forma ambigua", "Pulsos distintos con igual traza e igual espectro: anchos FWHM que "
            "difieren ≈ 25 % (Chung y Weiner, p. 665)."),
           ("Artefacto coherente", "Con pulsos inestables, un pico angosto subestima la duración; "
            "la real la da el pedestal (Ratner et al., p. 2874).")]
    for i, (t_, b_) in enumerate(lim):
        y = TOP + i * 1.72
        rect(s, 7.05, y, 5.68, 1.55, RED_T if i == 2 else GRAY_T)
        text(s, 7.3, y + 0.15, 5.2, 0.4, t_, 17, RED if i == 2 else VIOLET, HEAD)
        text(s, 7.3, y + 0.58, 5.2, 0.95, b_, 14)

    # 17 · Ishikawa --------------------------------------------------------------------
    s = d.content("5 · ¿Qué puede cambiar el ancho medido?",
                  "¿Cuáles son las fuentes de variación del ancho medido?",
                  "Diagrama de causa y efecto con todo lo que puede mover el ancho medido, en seis "
                  "ramas. La lectura clave: el ancho medido cambia por dos razones distintas, "
                  "porque el pulso que llega cambió (fuente y propagación) o porque la medición lo "
                  "deforma (proceso no lineal, geometría, detección e interpretación). Las dos "
                  "diapositivas siguientes detallan cada grupo.")
    image(s, F("ishikawa"), (SW - 10.8) / 2, TOP - 0.15, w=10.8)
    rect(s, MX, 6.2, 12.13, 0.7, VIO_T)
    text(s, MX + 0.25, 6.2, 11.6, 0.7, "El ancho medido cambia porque **el pulso cambia** o "
         "porque **la medición lo deforma**.", 18, anchor=MSO_ANCHOR.MIDDLE)

    # 18 · Factores del pulso -------------------------------------------------------------
    s = d.content("5 · ¿Qué puede cambiar el ancho medido?", "¿Qué factores cambian el pulso que llega?",
                  "Primer grupo: lo que cambia el pulso real. En la fuente, el bombeo, la alineación "
                  "de la cavidad, su compensación interna de dispersión y el calentamiento; Sala et "
                  "al. muestran que desajustar la cavidad 5 µm deshace o ensancha el pulso (p. 995). "
                  "En la propagación, el vidrio atravesado y cualquier compresor externo. Con "
                  "potencia alta aparece la automodulación de fase, que cambia el espectro.",
                  "Sala et al. (1980), p. 995; Diels y Rudolph (2006), cap. 1")
    text(s, MX, TOP, 12.1, 0.45, "Grupo 1 · Cambian **el pulso real**", 18, RED, HEAD)
    groups = [("Fuente", ["Potencia de bombeo", "Alineación de la cavidad",
                          "Compensación interna de dispersión", "Temperatura y calentamiento"]),
              ("Propagación", ["Material y espesor de vidrio (GDD)", "Compresor externo de "
                               "prismas o rejillas"]),
              ("Potencia alta", ["Automodulación de fase (SPM): cambia el espectro y, con él, el "
                                 "pulso"])]
    for i, (t_, items) in enumerate(groups):
        card(s, MX + i * 4.1, TOP + 0.65, 3.85, 3.0, t_, items, size=15)
    quote(s, MX, TOP + 3.95, 12.13, 1.1, "… a shortening by as little as 5 μm producing virtually "
          "complete disintegration of the dye laser output", "Sala et al. (1980), p. 995", 16)

    # 19 · Factores de la medición ------------------------------------------------------------
    s = d.content("5 · ¿Qué puede cambiar el ancho medido?", "¿Qué factores deforman la medición?",
                  "Segundo grupo: lo que deforma la medición aunque el pulso no cambie. Cristal: un "
                  "espesor excesivo alarga la traza (Diels y Rudolph, p. 468). Geometría: ángulo de "
                  "cruce, tamaño del haz, superposición de los brazos. Retardo: calibración y "
                  "linealidad del barrido. Detección: saturación, ruido y el promedio sobre pulsos "
                  "que pueden no ser idénticos (pp. 465–466). Interpretación: forma supuesta, método "
                  "de ajuste y nivel del fondo.",
                  "Diels y Rudolph (2006), pp. 465–468")
    text(s, MX, TOP, 12.1, 0.45, "Grupo 2 · Deforman **la medición**", 18, VIOLET, HEAD)
    rows = [("Cristal", "Espesor y desajuste de velocidad de grupo (alarga la traza); ángulo de "
             "ajuste de fase."),
            ("Geometría", "Ángulo de cruce y tamaño del haz; superposición de los brazos; enfoque."),
            ("Retardo", "Calibración del motor; velocidad de barrido frente al tiempo de "
             "integración."),
            ("Detección", "Saturación o no linealidad del detector; ruido; pulsos promediados "
             "que pueden no ser idénticos."),
            ("Interpretación", "Forma supuesta; método de ajuste; nivel del fondo.")]
    for i, (t_, b_) in enumerate(rows):
        y = TOP + 0.6 + i * 0.9
        rect(s, MX, y, 12.13, 0.78, GRAY_T if i % 2 == 0 else WHITE)
        text(s, MX + 0.25, y, 2.6, 0.78, t_, 17, VIOLET, HEAD, anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + 2.9, y, 9.0, 0.78, b_, 16, anchor=MSO_ANCHOR.MIDDLE)

    # 20 · Qué hacemos con los factores -------------------------------------------------------
    s = d.content("5 · ¿Qué puede cambiar el ancho medido?", "¿Qué hacemos con todos estos factores?",
                  "Esta es la bisagra del argumento. Los factores son muchos, interactúan y algunos "
                  "no se pueden separar en la traza (el chirp de la fuente y el GDD del montaje "
                  "producen el mismo efecto). Si los variamos ahora, cualquier efecto se confunde con "
                  "un error del instrumento que todavía no conocemos. Por eso en este trabajo se "
                  "fijan, cada uno con una razón. Lo que no se puede fijar es el error "
                  "experimental, y eso es lo que estiman las repeticiones de la comparación.")
    table(s, MX, TOP, 8.1, [
        ["Factor", "Cómo se fija", "Por qué"],
        ["Forma supuesta", "La misma del fabricante", "Cambia el ancho hasta un 9 %"],
        ["Vidrio en el camino", "Medir cerca de la salida; descontar el GDD", "Alarga el pulso"],
        ["Calentamiento", "Mismo tiempo antes de cada serie", "Estabiliza el láser"],
        ["Potencia de entrada", "Fija, sin saturar el detector", "Evita deformar la traza"],
        ["Eje de retardo", "Calibrado antes de comparar", "Su error pasa completo al ancho"],
        ["Alineación", "Relación pico/fondo verificada", "Detecta un montaje desajustado"],
    ], [2.2, 3.1, 2.8], 14, row_h=0.6)
    rect(s, 9.0, TOP, 3.73, 2.6, RED_T)
    text(s, 9.25, TOP + 0.2, 3.3, 0.4, "DECISIÓN", 13, RED, HEAD)
    text(s, 9.25, TOP + 0.65, 3.3, 1.8, "En este trabajo los factores **se fijan**; no se "
         "varían.", 22, INK, HEAD)
    card(s, 9.0, TOP + 2.85, 3.73, 2.2, "Lo que no se puede fijar", "Fluctuaciones del láser y "
         "ruido: son el **error experimental** que estiman las repeticiones.", size=15)

    # 21 · Por qué calibrar --------------------------------------------------------------------
    s = d.content("6 · ¿Por qué empezar por calibrar?", "¿Por qué no estudiar directamente los factores?",
                  "Cuatro razones. Un estudio de factores mide cambios del ancho; si el instrumento "
                  "tiene un sesgo desconocido, esos cambios no se pueden interpretar. El "
                  "autocorrelador no se puede verificar con otro reloj, así que la referencia "
                  "disponible es el fabricante. La comparación responde una pregunta práctica sobre "
                  "el láser. Y es factible en el semestre. La literatura respalda este uso: "
                  "Monmayrant et al. recomiendan la autocorrelación como chequeo diario e "
                  "independiente (p. 28), y Diels y Rudolph exigen conocer la función de "
                  "transferencia del instrumento (p. 457).",
                  "Diels y Rudolph (2006), p. 457; Monmayrant et al. (2010), p. 28")
    args = [("Sin calibrar no hay interpretación", "Un estudio de factores mide **cambios**; con "
             "un sesgo desconocido no se pueden atribuir."),
            ("La única referencia disponible", "El autocorrelador no se verifica con un reloj: el "
             "fabricante midió con una técnica equivalente."),
            ("Una pregunta práctica", "¿El láser sigue entregando lo que promete?"),
            ("Factible en el semestre", "Pocos recursos; el estudio de factores se apoya en ella.")]
    for i, (t_, b_) in enumerate(args):
        x = MX + (i % 2) * 6.23
        y = TOP + (i // 2) * 1.6
        card(s, x, y, 5.9, 1.42, t_, b_, size=15, tsize=17)
    quote(s, MX, TOP + 3.3, 12.13, 1.7, "Any measurement technique introduces some perturbation "
          "on the parameter to be measured… A careful analysis of the diagnostic instrument is "
          "required to find its exact transfer function.", "Diels y Rudolph (2006), p. 457", 16)

    # 22 · Qué se compara ------------------------------------------------------------------------
    s = d.content("6 · ¿Por qué empezar por calibrar?", "¿Qué se compara exactamente?",
                  "Dos lados. Patrón: el ancho máximo del fabricante, < 90 fs, con su ancho espectral "
                  "de 8 a 10 nm y su tasa de repetición. Técnica: el ancho medido con el autocorrelador, con la "
                  "misma forma supuesta y las condiciones fijadas. La respuesta es τp en fs; la "
                  "respuesta secundaria es el ancho espectral, si hay espectrómetro, para calcular "
                  "Δν·τp. Las repeticiones deben ser independientes: realinear o volver a barrer "
                  "entre una y otra, y repartirlas en días distintos para capturar la variación "
                  "real del láser y del montaje.")
    rect(s, MX, TOP, 5.45, 2.75, RED_T)
    text(s, MX + 0.3, TOP + 0.22, 4.9, 0.4, "PATRÓN", 14, RED, HEAD)
    text(s, MX + 0.3, TOP + 0.65, 4.9, 0.6, "Dato del fabricante", 22, INK, HEAD)
    bullets(s, MX + 0.3, TOP + 1.35, 4.9, 1.5, ["Ancho de pulso máximo: **< 90 fs** (sech², "
                                                "factor 0,65)", "Ancho espectral 8 – 10 nm; "
                                                "f_{rep} 50 u 80 MHz"], 16)
    rect(s, 6.08, TOP + 1.0, 1.15, 1.15, INK, shape=MSO_SHAPE.OVAL)
    text(s, 6.08, TOP + 1.0, 1.15, 1.15, "vs", 24, WHITE, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    rect(s, 7.28, TOP, 5.45, 2.75, VIO_T)
    text(s, 7.58, TOP + 0.22, 4.9, 0.4, "TÉCNICA", 14, VIOLET, HEAD)
    text(s, 7.58, TOP + 0.65, 4.9, 0.6, "Autocorrelador de intensidad", 22, INK, HEAD)
    bullets(s, 7.58, TOP + 1.35, 4.9, 1.5, ["Misma forma supuesta que el fabricante",
                                            "Condiciones fijadas (diapositiva 20)"], 16,
            marker=VIOLET)
    items = [("Variable respuesta", "τ_{p} (fs), ancho de pulso deconvolucionado"),
             ("Respuesta secundaria", "Δλ medido → Δν·τ_{p}, si hay espectrómetro"),
             ("Repeticiones", "n mediciones independientes, en días distintos")]
    for i, (t_, b_) in enumerate(items):
        card(s, MX + i * 4.1, TOP + 3.0, 3.85, 1.5, t_, b_, size=15, tsize=16)

    # 23 · Hipótesis --------------------------------------------------------------------------------
    s = d.content("6 · ¿Por qué empezar por calibrar?", "¿Cuál es la hipótesis del experimento?",
                  "Es la comparación de una media con un valor especificado (Montgomery, cap. 2, "
                  "sec. 2-4.6): la media de los anchos medidos frente al ancho del fabricante. "
                  "Montgomery menciona las especificaciones como origen válido del valor de "
                  "referencia. Como la hoja de datos da un máximo (< 90 fs) y no un valor nominal, "
                  "la prueba es de una cola: solo un pulso más largo que 90 fs contradice al "
                  "fabricante. Si el "
                  "laboratorio tiene otro autocorrelador de referencia, la comparación pasa a ser "
                  "de dos tratamientos o pareada (secs. 2-4 y 2-5): es una pregunta para la "
                  "asesoría.", "Montgomery (2004), cap. 2, sec. 2-4.6, ec. 2-37")
    rect(s, MX, TOP, 7.1, 1.35, RED_T)
    text(s, MX + 0.3, TOP, 1.2, 1.35, "H₀:", 30, RED, HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + 1.5, TOP, 5.4, 1.35, ["μ_{τ} = 90 fs", "el ancho medido está en el límite que garantiza el "
                                       "fabricante: láser e instrumento en especificación"], 17, INK,
         anchor=MSO_ANCHOR.MIDDLE, space=2)
    for r in s.shapes[-1].text_frame.paragraphs[0].runs:
        r.font.size = Pt(26)
    rect(s, MX, TOP + 1.55, 7.1, 1.35, VIO_T)
    text(s, MX + 0.3, TOP + 1.55, 1.2, 1.35, "H_{A}:", 30, VIOLET, HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + 1.5, TOP + 1.55, 5.4, 1.35, ["μ_{τ} > 90 fs", "el pulso es más largo que lo que "
                                              "garantiza el fabricante"], 17, INK, anchor=MSO_ANCHOR.MIDDLE,
         space=2)
    for r in s.shapes[-1].text_frame.paragraphs[0].runs:
        r.font.size = Pt(26)
    equation(s, 8.0, TOP, 4.73, 2.9, "t₀ = (ȳ − τ_{fab}) / (S / √n)", 24,
             caption="Montgomery, ec. 2-37; una cola: se rechaza H₀ si t₀ > t_{α, n−1}")
    card(s, MX, TOP + 3.2, 5.9, 1.65, "Por qué es de una cola",
         "La hoja de datos garantiza un máximo (\"< 90 fs\"), no un valor nominal: un pulso más "
         "corto no contradice al fabricante.", size=15)
    card(s, 6.83, TOP + 3.2, 5.9, 1.65, "Pregunta para la asesoría",
         "Si hay otro autocorrelador de referencia: dos tratamientos o comparación pareada "
         "(secs. 2-4 y 2-5).", RED_T, RED, size=15)

    # 24 · Diagnóstico ------------------------------------------------------------------------------
    s = d.content("6 · ¿Por qué empezar por calibrar?", "¿Qué podremos concluir sobre el láser?",
                  "Árbol de diagnóstico. Si el ancho coincide con el fabricante y el producto "
                  "Δν·τp está cerca del límite, láser e instrumento están en especificación y la "
                  "técnica queda calibrada. Si el pulso sale más largo con el espectro nominal, "
                  "tiene chirp: dispersión sin compensar. Si sale más largo con el espectro más "
                  "angosto, el láser perdió modos en fase. Si sale más corto que el nominal, hay que "
                  "sospechar del instrumento. El espectro es lo que separa 'el láser cambió' de 'el "
                  "instrumento mide mal'; Donnelly y Grossman hacen esa misma comparación (pp. "
                  "682–683). Limitación a declarar: aun con espectro, la forma no queda fija.",
                  "Donnelly y Grossman (1998), pp. 682–683; Chung y Weiner (2001)")
    rect(s, 3.65, TOP, 6.0, 0.9, INK)
    text(s, 3.65, TOP, 6.0, 0.9, "Comparar τ_{p} con el fabricante y Δλ con el nominal", 17, WHITE,
         HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    leaves = [("τ_{p} ≈ fabricante", "Δν·τ_{p} cerca del límite", "Láser e instrumento en "
               "especificación: **técnica calibrada**", VIO_T, VIOLET),
              ("τ_{p} mayor", "espectro igual al nominal", "El pulso tiene **chirp**: dispersión "
               "sin compensar", GRAY_T, INK),
              ("τ_{p} mayor", "espectro más angosto", "Menos modos en fase: el **láser perdió "
               "desempeño**", RED_T, RED),
              ("τ_{p} menor", "que el nominal", "Sospechar del **instrumento**: retardo, forma "
               "supuesta o artefacto", GRAY_T, INK)]
    for i, (a, b, c, fill, col) in enumerate(leaves):
        x = MX + i * 3.07
        arrow(s, 6.65, TOP + 0.92, x + 1.42, TOP + 1.38, MUTED, 1.5)
        rect(s, x, TOP + 1.4, 2.85, 3.0, fill)
        text(s, x + 0.2, TOP + 1.55, 2.45, 0.45, a, 18, col, HEAD)
        text(s, x + 0.2, TOP + 2.0, 2.45, 0.45, b, 14, MUTED)
        text(s, x + 0.2, TOP + 2.55, 2.45, 1.8, c, 16)
    text(s, MX, TOP + 4.6, 12.1, 0.5, "El espectro separa \"el láser cambió\" de \"el instrumento "
         "mide mal\". Límite: aun con espectro, la forma del pulso no queda fija (≈ 25 %).", 15,
         MUTED)

    # 25 · Qué sigue ----------------------------------------------------------------------------------
    s = d.dark("Cerrar con los pasos siguientes y la pregunta para la asesoría. El diseño factorial "
               "sobre los factores del bloque 5 queda como trabajo futuro, sobre un instrumento ya "
               "calibrado.")
    text(s, MX, 0.9, 10, 0.4, "6 · ¿POR QUÉ EMPEZAR POR CALIBRAR?", 13, VIOLET_L, HEAD)
    text(s, MX, 1.3, 11, 0.9, "¿Qué sigue?", 40, WHITE, HEAD)
    nxt = [("1", "Hoja de pruebas del equipo", "La ficha ya fija el patrón (< 90 fs, sech²); falta "
            "la medición de fábrica de este láser."),
           ("2", "Calibrar el retardo y alinear", "Escala de τ verificada; relación pico/fondo; "
            "sin pedestal."),
           ("3", "Serie de mediciones repetidas", "Comparación simple frente al fabricante."),
           ("4", "Trabajo futuro", "Diseño factorial sobre los factores, con el instrumento ya "
            "calibrado.")]
    for i, (n_, t_, b_) in enumerate(nxt):
        x = MX + i * 3.07
        rect(s, x, 2.75, 2.85, 2.9, RGBColor(0x16, 0x1D, 0x33))
        text(s, x + 0.25, 2.95, 2.4, 0.7, n_, 34, RED if i < 3 else VIOLET_L, HEAD)
        text(s, x + 0.25, 3.7, 2.4, 0.8, t_, 18, WHITE, HEAD)
        text(s, x + 0.25, 4.5, 2.4, 1.1, b_, 14.5, RGBColor(0xC9, 0xD1, 0xDE))
    text(s, MX, 6.05, 12.1, 0.6, "Pregunta para la asesoría: ¿basta el dato del fabricante como "
         "patrón, o hay un instrumento de referencia en el laboratorio?", 16,
         RGBColor(0xC9, 0xD1, 0xDE), LIGHT, italic=True)

    # 26 · Referencias ---------------------------------------------------------------------------------
    s = d.content("Referencias", "Referencias",
                  "Referencias citadas en las diapositivas, en Normas ICONTEC (NTC 5613). La lista "
                  "completa de la búsqueda está en bibliografia/autocorrelador-intensidad-"
                  "referencias.txt.")
    refs = [
        "ARMSTRONG, J. A. Measurement of picosecond laser pulse widths. En: Applied Physics "
        "Letters. 1967. vol. 10, no. 1, p. 16-18.",
        "CHUNG, Jung-Ho y WEINER, Andrew M. Ambiguity of ultrashort pulse shapes retrieved from "
        "the intensity autocorrelation and the power spectrum. En: IEEE Journal of Selected "
        "Topics in Quantum Electronics. 2001. vol. 7, no. 4, p. 656-666.",
        "DIELS, Jean-Claude y RUDOLPH, Wolfgang. Ultrashort laser pulse phenomena: fundamentals, "
        "techniques, and applications on a femtosecond time scale. 2 ed. Burlington: Academic "
        "Press, 2006.",
        "DONNELLY, T. D. y GROSSMAN, Carl. Ultrafast phenomena: a laboratory experiment for "
        "undergraduates. En: American Journal of Physics. 1998. vol. 66, no. 8, p. 677-685.",
        "HAUS, Hermann A. Theory of mode locking with a fast saturable absorber. En: Journal of "
        "Applied Physics. 1975. vol. 46, no. 7, p. 3049-3058.",
        "MONMAYRANT, Antoine; WEBER, Sébastien y CHATEL, Béatrice. A newcomer's guide to "
        "ultrashort pulse shaping and characterization. En: Journal of Physics B: Atomic, "
        "Molecular and Optical Physics. 2010. vol. 43, no. 10, art. 103001.",
        "MONTGOMERY, Douglas C. Diseño y análisis de experimentos. 2 ed. México: Limusa Wiley, "
        "2004.",
        "RATNER, Justin, et al. Coherent artifact in modern pulse measurements. En: Optics "
        "Letters. 2012. vol. 37, no. 14, p. 2874-2876.",
        "SALA, K.; KENNEY-WALLACE, G. y HALL, G. CW autocorrelation measurements of picosecond "
        "laser pulses. En: IEEE Journal of Quantum Electronics. 1980. vol. 16, no. 9, p. 990-996.",
    ]
    tf = box(s, MX, TOP - 0.1, 12.13, 5.3)
    for i, r in enumerate(refs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.4))))
        pPr.set("indent", str(-int(Inches(0.4))))
        runs(p, r, 14, INK)

    d.prs.save(out)
    print("pptx:", out, "·", d.n, "diapositivas")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
