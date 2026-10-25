"""Análisis independiente del piloto del 24 oct 2026 (trazas R01 y R08–R15).

Lee las trazas del libro usado, reproduce el cálculo de ancho a media altura del Excel, ajusta por
mínimos cuadrados una autocorrelación sech² y una gaussiana a cada traza, resume por corriente y
calcula el número de réplicas con la t no central exacta (scipy). Escribe un JSON con todo.

Uso: .venv/bin/python scripts/trabajo-practico/analiza_piloto.py <libro.xlsx> <salida.json>
"""
import json
import sys

import numpy as np
import openpyxl
from scipy import optimize, stats

VALIDAS = ["R01"] + [f"R{k:02d}" for k in range(8, 16)]
CONV = 2 * 1e-6 / 299792458.0 * 1e15          # fs por µm (2 pasos)
F_SECH, F_GAUSS = 0.65, 0.7071                 # factores de forma τp = F · Δτ_AC
K_SECH = 1.7627                                # τp = K · t0 para sech²(t/t0)
TAU0 = {"I1": 107.0, "I2": 76.0, "I3": 54.0}
ALFA, POT, DELTA, CONF = 0.05, 0.90, 10.0, 0.80


def g_sech2(tau, tau_p):
    u = tau / (tau_p / K_SECH)
    u = np.where(np.abs(u) < 1e-6, 1e-6, u)
    return 3 * (u / np.tanh(u) - 1) / np.sinh(u) ** 2


def g_gauss(tau, tau_p):
    d_ac = tau_p / F_GAUSS                     # FWHM de la autocorrelación
    return np.exp(-4 * np.log(2) * tau ** 2 / d_ac ** 2)


def fwhm_excel(tau, y, n_fondo=2):
    """Mismo algoritmo que el libro: fondo de las alas, normalizar, cruces a 0,5 interpolados."""
    f = (y[:n_fondo].mean() + y[-n_fondo:].mean()) / 2
    e = (y - f) / (y.max() - f)
    p = int(np.argmax(e))
    n = len(e)
    il = max(i for i in range(p) if e[i] < 0.5 <= e[i + 1])
    ir = min(i for i in range(p, n - 1) if e[i] >= 0.5 > e[i + 1])
    cr = lambda i: tau[i] + (0.5 - e[i]) * (tau[i + 1] - tau[i]) / (e[i + 1] - e[i])
    return cr(ir) - cr(il), (cr(ir) + cr(il)) / 2, int((e >= 0.5).sum())


def ajuste(tau, y, g, tau_p0, centro0):
    modelo = lambda t, A, B, tc, tp: B + A * g(t - tc, tp)
    A0, B0 = y.max() - y.min(), y.min()
    p, cov = optimize.curve_fit(modelo, tau, y, p0=[A0, B0, centro0, tau_p0], maxfev=20000)
    res = y - modelo(tau, *p)
    r2 = 1 - (res ** 2).sum() / ((y - y.mean()) ** 2).sum()
    err = np.sqrt(np.diag(cov))
    return {"A": p[0], "B": p[1], "centro": p[2], "tau_p": abs(p[3]), "se_tau_p": err[3], "r2": r2,
            "sd_res_rel": res.std(ddof=4) / p[0]}


def potencia(n, d, alfa=ALFA, colas=2):
    nu = n - 1
    lam = d * np.sqrt(n)
    if colas == 2:
        t = stats.t.ppf(1 - alfa / 2, nu)
        p = stats.nct.sf(t, nu, lam) + stats.nct.cdf(-t, nu, lam)
    else:
        p = stats.nct.sf(stats.t.ppf(1 - alfa, nu), nu, lam)
    return 1.0 if np.isnan(p) else float(p)          # scipy devuelve NaN con λ grande (potencia ≈ 1)


def n_necesario(d, pot=POT):
    return next((n for n in range(2, 3001) if potencia(n, d) >= pot), None)


def delta_det(n, sigma, pot=POT):
    if potencia(n, 6.0) < pot:
        return None
    return optimize.brentq(lambda d: potencia(n, d) - pot, 1e-3, 6.0) * sigma


def main(libro, salida):
    wb = openpyxl.load_workbook(libro, data_only=True)
    reg = wb["Registro piloto"]
    trazas = {}
    for r in range(5, 20):
        hoja = reg.cell(r, 2).value
        trazas[hoja] = {"n": reg.cell(r, 1).value, "corriente": reg.cell(r, 3).value, "I_A": reg.cell(r, 4).value,
                        "inicio": reg.cell(r, 5).value, "fin": reg.cell(r, 6).value,
                        "fwhm_excel": reg.cell(r, 10).value, "taup_excel": reg.cell(r, 11).value,
                        "estado_excel": reg.cell(r, 12).value}
    res = {}
    for hoja in VALIDAS:
        ws = wb[hoja]
        x, y = [], []
        r = 2
        while ws.cell(r, 1).value is not None:
            x.append(float(ws.cell(r, 1).value))
            y.append(float(ws.cell(r, 2).value))
            r += 1
        x, y = np.array(x), np.array(y)
        tau = x * CONV
        d_ac, centro, k_in = fwhm_excel(tau, y)
        t = dict(trazas[hoja])
        t.update({"puntos": len(x), "paso_um": float(np.median(np.diff(x))), "paso_fs": float(np.median(np.diff(tau))),
                  "x": x.tolist(), "y": y.tolist(),
                  "fwhm_py": d_ac, "lect_fwhm": k_in, "taup_fwhm_sech": F_SECH * d_ac, "taup_fwhm_gauss": F_GAUSS * d_ac,
                  "fit_sech": ajuste(tau, y, g_sech2, F_SECH * d_ac, centro),
                  "fit_gauss": ajuste(tau, y, g_gauss, F_GAUSS * d_ac, centro)})
        t["pico_uW"] = float(y.max())
        res[hoja] = t

    estim = {"fwhm_sech": lambda t: t["taup_fwhm_sech"], "fwhm_gauss": lambda t: t["taup_fwhm_gauss"],
             "fit_sech": lambda t: t["fit_sech"]["tau_p"], "fit_gauss": lambda t: t["fit_gauss"]["tau_p"]}
    resumen = {}
    for variante, excluir in (("con_R01", ()), ("sin_R01", ("R01",))):
      for nombre, f in estim.items():
          por = {}
          for cur in ("I1", "I2", "I3"):
              v = np.array([f(t) for h, t in res.items() if t["corriente"] == cur and h not in excluir])
              por[cur] = {"n": int(v.size), "media": float(v.mean()), "S": float(v.std(ddof=1)) if v.size > 1 else None,
                          "valores": v.tolist(), "tau0": TAU0[cur], "dif": float(v.mean() - TAU0[cur])}
          # réplicas: σ conservadora de la mayor S
          con_s = {c: p for c, p in por.items() if p["S"] is not None}
          cmax = max(con_s, key=lambda c: con_s[c]["S"])
          nS = con_s[cmax]["n"]
          S = con_s[cmax]["S"]
          fac = float(np.sqrt((nS - 1) / stats.chi2.ppf(1 - CONF, nS - 1)))
          sigma = fac * S
          d = DELTA / sigma
          n = n_necesario(d)
          # mismo cálculo solo con I1 (5 trazas, las únicas con n = 5)
          S1 = por["I1"]["S"]
          fac1 = float(np.sqrt(4 / stats.chi2.ppf(1 - CONF, 4)))
          n1 = n_necesario(DELTA / (fac1 * S1))
          resumen[f"{nombre}|{variante}"] = {"por_corriente": por, "S_max": S, "corriente_S_max": cmax, "n_S_max": nS, "factor": fac,
                             "sigma": sigma, "d": d, "n": n, "potencia_n": float(potencia(n, d)) if n else None,
                             "sigma_I1": fac1 * S1, "n_I1": n1,
                             "delta_det": {k: delta_det(k, sigma) for k in range(2, 31)}}

    # tiempos
    def mins(a, b):
        return ((b.hour * 60 + b.minute) - (a.hour * 60 + a.minute))
    dur = [mins(t["inicio"], t["fin"]) for h, t in res.items() if t["inicio"] and t["fin"]]
    inicios = [res[h]["inicio"] for h in VALIDAS[1:]]                # R08–R15 consecutivas
    entre = [mins(a, b) for a, b in zip(inicios[:-1], inicios[1:])]
    tiempos = {"toma_sola_min": dur, "toma_sola_media": float(np.mean(dur)),
               "entre_inicios_R08_R15": entre, "entre_inicios_media": float(np.mean(entre)),
               "todas_15_entre_inicios": trazas and float(mins(trazas["R01"]["inicio"], trazas["R15"]["inicio"]) / 14)}
    json.dump({"trazas": res, "resumen": resumen, "tiempos": tiempos, "registro": {
        h: {k: (str(v) if v is not None else None) for k, v in t.items() if k in ("corriente", "inicio", "fin", "estado_excel", "taup_excel")}
        for h, t in trazas.items()}}, open(salida, "w"), indent=1, default=str)
    for nombre, s in resumen.items():
        print(nombre, {c: (p["n"], round(p["media"], 1), round(p["S"], 2) if p["S"] else None) for c, p in s["por_corriente"].items()},
              "Smax", round(s["S_max"], 2), s["corriente_S_max"], "sigma", round(s["sigma"], 2), "n", s["n"], "n_I1", s["n_I1"])
    print("tiempos", tiempos)
    for h, t in res.items():
        print(h, t["corriente"], t["puntos"], round(t["paso_fs"], 1), "excel", round(t["fwhm_excel"], 2), "py", round(t["fwhm_py"], 2),
              "sech", round(t["taup_fwhm_sech"], 1), round(t["fit_sech"]["tau_p"], 1), "gauss", round(t["taup_fwhm_gauss"], 1),
              round(t["fit_gauss"]["tau_p"], 1), "r2", round(t["fit_sech"]["r2"], 3), round(t["fit_gauss"]["r2"], 3), "lect", t["lect_fwhm"])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
