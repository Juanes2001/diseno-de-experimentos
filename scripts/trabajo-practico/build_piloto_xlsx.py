"""Genera el libro de Excel del piloto del sábado 24 oct 2026 (trabajo práctico).

El libro recibe, por cada traza del autocorrelador, el desplazamiento del motor y la potencia
óptica medida; calcula con fórmulas el retardo, la traza normalizada, el ancho a media altura y
τp; resume el piloto por corriente y calcula el número de réplicas de la semana con el método del
protocolo (σ = límite superior de confianza de S, δ por la regla de precedencia, curva OC de la t
de una muestra) y el plan de medición por día.

Uso:
    .venv/bin/python scripts/trabajo-practico/build_piloto_xlsx.py            # libro final
    .venv/bin/python scripts/trabajo-practico/build_piloto_xlsx.py --prueba X  # copia pequeña con
                                                                               # datos sintéticos
"""
import argparse
import random
from pathlib import Path

import numpy as np
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell as rc

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = RAIZ / "trabajo-practico" / "piloto" / "Piloto_Replicas_Autocorrelador.xlsx"

N_TRAZAS = 15            # 5 por corriente
ESTIMADORES = ["FWHM sech²", "FWHM gauss", "Ajuste sech²", "Ajuste gauss"]

# Hoja de pruebas de fábrica del láser (serie 881139), tabla 5 «Laser Output Characterization»
# (fotos del 24 oct 2026 en notas/). Se eligen las tres corrientes altas alcanzables (máx. 7,12 A).
CORRIENTES = (5.0, 6.0, 7.0)           # A (AMP pump current)
TAU0_FAB = (107, 76, 54)               # fs, ancho de pulso de la tabla 5
P_FAB = (218, 378, 579)                # mW, potencia media de la tabla 5
TAUMAX_FAB = ("", "", 80)              # fs, especificación «< 80» de la tabla 4 (a 7,0 A)
SEMILLA_PILOTO = 20261024
MAX_RPD = 20             # máximo de réplicas por corriente y por día en el plan
DIAS_MAX = 4
FONT = "Arial"

# ------------------------------------------------------------------ orden sorteado del piloto


def orden_piloto():
    lista = ["I1"] * 5 + ["I2"] * 5 + ["I3"] * 5
    random.Random(SEMILLA_PILOTO).shuffle(lista)
    return lista


# ------------------------------------------------------------------ traza sintética (ejemplo)


FS_POR_UM = 2 * 1e-6 / 299792458.0 * 1e15                   # retardo por µm de micrómetro (2 pasos)


def g_sech2(tau, tau_p_fs):
    """Autocorrelación de intensidad normalizada de un pulso sech² de ancho τp (sin fondo coherente)."""
    u = np.where(np.abs(tau) < 1e-6, 1.0, tau / (tau_p_fs / 1.7627))
    return np.where(np.abs(tau) < 1e-6, 1.0, 3 * (u / np.tanh(u) - 1) / np.sinh(u) ** 2)


def traza_manual(tau_p_fs, n_zona=16, colas=(2.0, 3.0), zona=1.5, resol_um=1.0, ruido=0.015, fondo=0.02,
                 x0_um=12500.0, semilla=1):
    """Traza tomada a mano con micrómetro, como la del montaje real (≈ 20 lecturas).

    n_zona lecturas con paso uniforme en la zona de solapamiento (± zona·Δτ_AC del pico) y una lectura
    por ala a cada distancia de «colas» (en múltiplos de Δτ_AC); posiciones redondeadas a la
    resolución del micrómetro, en µm; potencia óptica en W.
    """
    rng = np.random.default_rng(semilla)
    ac_um = tau_p_fs / 0.65 / FS_POR_UM
    paso = max(resol_um, np.floor(2 * zona * ac_um / (n_zona - 1) / resol_um) * resol_um)
    desfase = rng.uniform(-paso / 2, paso / 2)               # el pico no cae justo en una lectura
    xz = (np.arange(n_zona) - (n_zona - 1) / 2) * paso
    xc = np.array(colas) * ac_um
    xs = np.concatenate([-xc[::-1], xz, xc]) + desfase
    x_um = np.unique(np.round(xs / resol_um) * resol_um)      # sin lecturas repetidas
    g = g_sech2(x_um * FS_POR_UM, tau_p_fs)
    y = (g + fondo) * (1 + ruido * rng.standard_normal(x_um.size))
    return x0_um + x_um, np.round(y * 3.5e-5, 9)             # posición del micrómetro (µm), potencia (W)


# ------------------------------------------------------------------ libro


def build(salida, max_filas, datos_prueba=None, registro=None, ajustes=None, nota_registro=None):
    wb = xlsxwriter.Workbook(str(salida))
    wb.set_calc_mode("auto")
    wb.set_properties({"title": "Piloto · réplicas del autocorrelador",
                       "author": "Juan Esteban Rodríguez Ochoa · Sebastián Zapata Henao"})

    base = {"font_name": FONT, "font_size": 10, "valign": "vcenter"}

    def F(**kw):
        d = dict(base)
        d.update(kw)
        return wb.add_format(d)

    f_tit = F(bold=True, font_size=14, font_color="#1F3864")
    f_sub = F(italic=True, font_color="#595959")
    f_h = F(bold=True, bg_color="#1F3864", font_color="white", border=1, text_wrap=True, align="center")
    f_lab = F(border=1, text_wrap=True)
    f_labb = F(border=1, bold=True, text_wrap=True)
    f_in = F(border=1, bg_color="#FFF2CC", font_color="#0000FF")
    f_in_num = F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="General")
    f_in_t = F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="hh:mm")
    f_in_txt = F(border=1, bg_color="#FFF2CC", font_color="#0000FF", text_wrap=True)
    f_calc = F(border=1, num_format="0.00")
    f_calc3 = F(border=1, num_format="0.000")
    f_sci = F(border=1, num_format="0.00E+00")
    f_int = F(border=1, num_format="0")
    f_link = F(border=1, font_color="#008000", num_format="0.00")
    f_linkt = F(border=1, font_color="#008000")
    f_key = F(border=2, bold=True, bg_color="#E2EFDA", num_format="0.00", font_size=12)
    f_keyi = F(border=2, bold=True, bg_color="#E2EFDA", num_format="0", font_size=12)
    f_keyt = F(border=2, bold=True, bg_color="#E2EFDA", text_wrap=True)
    f_txt = F(text_wrap=True, valign="top")
    f_gray = F(font_color="#808080", num_format="0.00")
    f_gray3 = F(font_color="#808080", num_format="0.0000")
    f_note = F(italic=True, font_color="#595959", text_wrap=True, valign="top")
    f_ok = F(bg_color="#C6EFCE", font_color="#006100", border=1)
    f_warn = F(bg_color="#FFEB9C", font_color="#9C5700", border=1)
    f_bad = F(bg_color="#FFC7CE", font_color="#9C0006", border=1)

    L = max_filas + 1                     # última fila (Excel, base 1) con datos

    # nombres de hoja
    S_CFG, S_REG, S_RES, S_REP, S_PLAN = "Configuración", "Registro piloto", "Resumen piloto", "Réplicas", "Plan semana"
    trazas = [f"R{k:02d}" for k in range(1, N_TRAZAS + 1)]
    q = lambda s: f"'{s}'"

    ws_ins = wb.add_worksheet("Instrucciones")
    ws_cfg = wb.add_worksheet(S_CFG)
    ws_reg = wb.add_worksheet(S_REG)
    ws_res = wb.add_worksheet(S_RES)
    ws_rep = wb.add_worksheet(S_REP)
    ws_plan = wb.add_worksheet(S_PLAN)
    S_AJ = "Ajustes"
    ws_aj = wb.add_worksheet(S_AJ)
    ws_tr = [wb.add_worksheet(n) for n in trazas]
    ws_ej = wb.add_worksheet("Ejemplo")

    # ================================================================ Configuración
    w = ws_cfg
    w.set_column("A:A", 46)
    w.set_column("B:B", 16)
    w.set_column("C:C", 70)
    w.write("A1", "Configuración del piloto y del cálculo de réplicas", f_tit)
    w.write("A2", "Celdas amarillas con letra azul = se pueden editar. Todo lo demás es fórmula.", f_sub)
    w.write_row("A4", ["Parámetro", "Valor", "Nota"], f_h)

    filas_cfg = {}

    def cfg(fila, nombre, etiqueta, valor, nota, fmt=None, formula=False):
        w.write(fila, 0, etiqueta, f_lab)
        if formula:
            w.write_formula(fila, 1, valor, fmt or f_calc3)
        else:
            w.write(fila, 1, valor, fmt or f_in)
        w.write(fila, 2, nota, f_note)
        wb.define_name(nombre, f"={q(S_CFG)}!$B${fila + 1}")
        filas_cfg[nombre] = fila

    r = 4
    w.write(r, 0, "1. Conversión del desplazamiento a retardo", f_labb); r += 1
    cfg(r, "unidad", "Unidad de la posición del micrómetro (columna A de R01–R15)", "µm",
        "Fija en µm (p. ej. 12483 para 12,483 mm). La potencia óptica de la columna B va en W, tal como la da el medidor.",
        F(border=1, bold=True)); r += 1
    cfg(r, "pasos", "Número de veces que el haz recorre el desplazamiento", 2,
        "2 = retrorreflector en el brazo móvil (ida y vuelta): τ = 2Δx/c. Usar 4 si el brazo es de doble paso.",
        f_in_num); r += 1
    cfg(r, "c_luz", "Velocidad de la luz c (m/s)", 299792458, "Valor exacto (SI).", F(border=1, num_format="0")); r += 1
    cfg(r, "a_metros", "Factor de la unidad a metros", '=IF(unidad="nm",10^-9,IF(unidad="mm",10^-3,10^-6))',
        "Automático según la unidad.", F(border=1, num_format="0E+00"), formula=True); r += 1
    cfg(r, "conv", "Retardo por unidad de desplazamiento (fs por unidad)", "=pasos*a_metros/c_luz*10^15",
        "Automático. Con µm y 2 pasos da 6,6713 fs por cada µm del micrómetro.", F(border=1, num_format="0.0000"), formula=True); r += 1
    r += 1
    w.write(r, 0, "2. Lectura del ancho de pulso", f_labb); r += 1
    cfg(r, "factor_forma", "Factor de forma: τp = factor × ancho de la traza", 0.65,
        "0,65 = el que usa Calmar (sech²; el valor teórico es 0,648). Gaussiano: 0,7071.", f_in_num); r += 1
    cfg(r, "estimador", "Estimador de τp que usan el Resumen, Réplicas y Plan semana", "Ajuste sech²",
        "FWHM = ancho a media altura leído en la traza × factor. Ajuste = mínimos cuadrados de la autocorrelación "
        "completa (hoja Ajustes; valores calculados con scipy). Con pasos gruesos el ajuste usa todos los puntos y es "
        "más estable; las cuatro opciones están en la hoja Ajustes."); r += 1
    w.data_validation(r - 1, 1, r - 1, 1, {"validate": "list", "source": ESTIMADORES})
    cfg(r, "n_fondo", "Lecturas en cada ala (fuera del solapamiento) para el fondo", 2,
        "El fondo es el promedio de las primeras y las últimas N lecturas de la traza. Con el montaje actual: 2 por "
        "ala, a 2 y 3 Δτ_AC del pico (ver la guía de muestreo).",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    cfg(r, "ventana", "Ventana de suavizado (puntos, número impar)", 1,
        "1 = sin suavizar. Con toma manual (pocos puntos) dejarla en 1: promediar lecturas vecinas ensancha la traza. "
        "No cambiarla durante el experimento: es parte del programa de análisis fijo del protocolo.",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    cfg(r, "semi", "Semiventana (automático)", "=INT((MAX(1,ventana)-1)/2)", "Puntos a cada lado del promedio móvil.",
        F(border=1, num_format="0"), formula=True); r += 1
    cfg(r, "pts_min", "Lecturas mínimas dentro del ancho a media altura", 5,
        "Por debajo de esto la interpolación de los cruces pierde precisión: afinar el paso cerca del pico.",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    cfg(r, "rec_min", "Recorrido mínimo a cada lado del centro (múltiplos de Δτ_AC)", 2,
        "Las lecturas del fondo deben quedar a 2 Δτ_AC o más del centro, donde la autocorrelación sech² es < 0,5 % del "
        "pico. Si quedan más cerca, el fondo sale alto y el ancho sale corto (≈ −1,5 fs en simulación).",
        f_in_num); r += 1
    cfg(r, "n_zona", "Lecturas en la zona de solapamiento", 16,
        "Con el montaje actual caben 16 o 17 lecturas en la zona donde hay solapamiento (señal). Más 2 por ala = 20.",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    cfg(r, "zona_ac", "Semiancho de la zona de solapamiento (múltiplos de Δτ_AC)", 1.5,
        "A ±1,5 Δτ_AC la señal es ≈ 1 % del pico: ahí termina en la práctica el solapamiento.", f_in_num); r += 1
    cfg(r, "resol", "Resolución de lectura del micrómetro (µm)", 1,
        "La menor división que se lee con seguridad (con nonio, 1 µm; sin nonio, 10 µm). Solo se usa en la guía de "
        "muestreo de abajo.", f_in_num); r += 1
    r += 1
    w.write(r, 0, "3. Corrientes del diseño", f_labb); r += 1
    for k in (1, 2, 3):
        cfg(r, f"I{k}_A", f"Corriente I{k} (A)", CORRIENTES[k - 1],
            "Corriente de bombeo del amplificador (AMP pump). Hoja de pruebas de fábrica, tabla 5. Se usan las tres "
            "corrientes altas alcanzables: el controlador llega a 7,12 A, así que la fila de 7,23 A queda fuera.",
            f_in_num); r += 1
    wb.define_name("corrientes", f"={q(S_CFG)}!$B${r - 2}:$B${r}")
    r += 1
    w.write(r, 0, "4. Número de réplicas (protocolo, sección 4)", f_labb); r += 1
    cfg(r, "alfa", "α (riesgo de error tipo I)", 0.05, "Protocolo: 0,05.", f_in_num); r += 1
    cfg(r, "potencia", "Potencia deseada 1 − β", 0.90, "Protocolo: 0,90.", f_in_num); r += 1
    cfg(r, "colas", "Colas de la prueba (1 o 2)", 2,
        "2 si τ0 es un valor medido por el fabricante (cualquier cambio importa); 1 si solo se tiene el máximo "
        "garantizado. Con 2 colas el n es algo mayor (más conservador).",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    w.data_validation(r - 1, 1, r - 1, 1, {"validate": "list", "source": [1, 2]})
    cfg(r, "conf_sigma", "Confianza del límite superior de σ", 0.80,
        "Protocolo: σ = límite superior del 80 % de S (con n = 5 multiplica S por 1,56).", f_in_num); r += 1
    cfg(r, "delta_def", "δ por defecto si el fabricante no da ningún dato (fs)", 10,
        "Protocolo, sección 4, regla 5.", f_in_num); r += 1
    r += 1
    w.write(r, 0, "5. Datos del fabricante para fijar δ (llenar los que existan; ver hoja Réplicas)", f_labb); r += 1
    fila_fab = r
    w.write_row(r, 0, ["Dato", "I1", "I2", "I3"], f_h)
    w.set_column("D:E", 12)
    r += 1
    datos_fab = [
        ("tau0", "τ0: ancho medido por el fabricante (fs)"),
        ("tol", "(1) Tolerancia del valor medido t, en τ0 ± t (fs)"),
        ("taumax", "(2) Máximo garantizado τmáx (fs)"),
        ("banda", "(3) Semiancho de la banda de aceptación (fs)"),
        ("estab", "(4) Estabilidad declarada del ancho (% rms, p. ej. 2 para 2 %)"),
        ("incert", "Incertidumbre u del fabricante (fs); piso δ ≥ 2u"),
    ]
    for clave, etiqueta in datos_fab:
        w.write(r, 0, etiqueta, f_lab)
        for k in (1, 2, 3):
            val = {"tau0": TAU0_FAB[k - 1], "taumax": TAUMAX_FAB[k - 1]}.get(clave, "")
            w.write(r, k, val, f_in_num)
            wb.define_name(f"{clave}_{k}", f"={q(S_CFG)}!${'BCD'[k - 1]}${r + 1}")
        r += 1
    w.write(r, 0, "Referencia: potencia media de fábrica (mW)", f_lab)
    for k in (1, 2, 3):
        w.write(r, k, P_FAB[k - 1], f_in_num)
        wb.define_name(f"pfab_{k}", f"={q(S_CFG)}!${'BCD'[k - 1]}${r + 1}")
    r += 1
    w.merge_range(r, 0, r, 2, "Fuente: hoja de pruebas de fábrica (tabla 5: 5 A → 107 fs, 6 A → 76 fs, 7 A → 54 fs). La "
            "hoja no da tolerancia, banda de aceptación, estabilidad del ancho ni incertidumbre. La estabilidad "
            "de 0,40 % rms es de la potencia, no del ancho, y por eso no se usa en la regla 4. "
            "τmáx = 80 fs es la especificación de la tabla 4, válida solo en el punto nominal (7,0 A). Ojo: la tabla 4 "
            "reporta 70 fs a 7,0 A y la tabla 5 reporta 54 fs a 7 A; se usa la tabla 5 porque da las tres "
            "corrientes con el mismo método.", f_note)
    w.set_row(r, 56)
    r += 2
    w.write(r, 0, "6. Tiempo disponible para la semana", f_labb); r += 1
    cfg(r, "horas_dia", "Horas de medición útiles por día (sin calentamiento)", 5,
        "Tiempo neto para tomar trazas cada día (lun 26 – jue 29).", f_in_num); r += 1
    cfg(r, "min_traza_man", "Minutos por traza (dejar vacío para usar el cronometraje del piloto)", "",
        "Si se llena, reemplaza la duración medida en el Registro del piloto.", f_in_num); r += 1
    cfg(r, "semilla", "Semilla del sorteo del orden de la semana", 2610,
        "Cambiarla genera otro orden aleatorio en la hoja Plan semana. Fijar el definitivo el sábado por la noche.",
        F(border=1, bg_color="#FFF2CC", font_color="#0000FF", num_format="0")); r += 1
    r += 1
    w.write(r, 0, "7. Guía de muestreo manual por corriente (automática)", f_labb); r += 1
    w.write_row(r, 0, ["Cálculo", "I1", "I2", "I3"], f_h); r += 1
    g0 = r + 1                                                   # fila Excel de la primera fila de la guía
    guia = [
        ("τ0 de fábrica (fs)", "=tau0_{k}", f_link),
        ("Ancho de la autocorrelación esperado Δτ_AC = τ0 / factor (fs)", "=IF(ISNUMBER({c}{f0}),{c}{f0}/factor_forma,\"\")", f_calc),
        ("Δτ_AC en recorrido del micrómetro (µm)", "=IF(ISNUMBER({c}{f1}),{c}{f1}/(pasos*10^-6/c_luz*10^15),\"\")", f_calc),
        ("Zona de solapamiento: ± desde el pico (µm)", "=IF(ISNUMBER({c}{f2}),ROUND(zona_ac*{c}{f2},0),\"\")", f_int),
        ("Paso en la zona = 2 × zona / (lecturas − 1) (µm)",
         "=IF(ISNUMBER({c}{f3}),MAX(resol,INT(2*zona_ac*{c}{f2}/(n_zona-1)/resol)*resol),\"\")", f_calc),
        ("Lecturas dentro del ancho a media altura", "=IF(ISNUMBER({c}{f4}),INT({c}{f2}/{c}{f4}),\"\")", f_int),
        ("Ala 1: lectura a ± 2 Δτ_AC del pico (µm)", "=IF(ISNUMBER({c}{f2}),ROUND(2*{c}{f2},0),\"\")", f_int),
        ("Ala 2: lectura a ± 3 Δτ_AC del pico (µm)", "=IF(ISNUMBER({c}{f2}),ROUND(3*{c}{f2},0),\"\")", f_int),
        ("Lecturas por traza = zona + 2 por ala", "=IF(ISNUMBER({c}{f2}),n_zona+2*n_fondo,\"\")", f_int),
        ("¿El micrómetro lo resuelve?", "=IF(ISNUMBER({c}{f5}),IF({c}{f5}>=pts_min,\"Sí\",\"No: hace falta leer más fino\"),\"\")", f_linkt),
    ]
    filas_g = {f"f{i}": g0 + i for i in range(len(guia))}
    for i, (et, fo, fm) in enumerate(guia):
        w.write(r, 0, et, f_lab)
        for k in (1, 2, 3):
            w.write_formula(r, k, fo.format(k=k, c="BCD"[k - 1], **filas_g), fm)
        r += 1
    w.merge_range(r, 0, r, 2, "Cómo tomar una traza a mano: buscar el pico; ir a la lectura del ala 2 de un lado; avanzar "
            "girando el micrómetro siempre en el mismo sentido (así el juego mecánico no cambia la posición): ala 2, ala 1, "
            "las lecturas de la zona con el paso de la tabla, ala 1 y ala 2 del otro lado. Anotar posición y potencia de "
            "cada lectura en orden.", f_note)
    w.set_row(r, 30)

    # ================================================================ hojas de traza
    def hoja_traza(w, nombre, num, datos=None):
        w.set_column("A:B", 15, f_in)
        w.set_column("C:E", 12)
        w.set_column("F:H", 9, None, {"hidden": False})
        w.set_column("I:J", 11)
        w.set_column("K:K", 7)
        w.set_column("L:L", 40)
        w.set_column("M:M", 14)
        w.freeze_panes(1, 0)
        w.write_row(0, 0, ["Posición del micrómetro (µm)", "Potencia óptica (W)", "Retardo τ (fs)",
                           "Señal suavizada", "Señal normalizada", "Cruce izq. (fila)", "Cruce der. (fila)",
                           "τ de cruce (fs)", "Retardo centrado (fs) · gráfica", "Normalizada · gráfica", "Sentido"], f_h)
        w.set_row(0, 42)
        rngA, rngB = f"$A$2:$A${L}", f"$B$2:$B${L}"
        rngC, rngD, rngE = f"$C$2:$C${L}", f"$D$2:$D${L}", f"$E$2:$E${L}"
        # panel de resultados (L:M)
        P = {}
        panel = [
            ("npts", "Puntos de la traza", f"=COUNT({rngA})", f_int),
            ("npts_b", "Puntos con potencia", f"=COUNT({rngB})", f_int),
            ("fondo", "Fondo (promedio de las alas)",
             f"=IF($M$3<2*n_fondo,\"\",(AVERAGE(INDEX({rngB},1):INDEX({rngB},n_fondo))"
             f"+AVERAGE(INDEX({rngB},$M$3-n_fondo+1):INDEX({rngB},$M$3)))/2)", f_sci),
            ("maxd", "Máximo de la señal suavizada", f"=IF($M$3<2*n_fondo,\"\",MAX({rngD}))", f_sci),
            ("filapico", "Fila del máximo", f"=IFERROR(MATCH($M$5,{rngD},0)+1,\"\")", f_int),
            ("taupico", "Retardo del máximo (fs)", f"=IFERROR(INDEX($C$1:$C${L},$M$6),\"\")", f_calc),
            ("filaizq", "Fila del cruce izquierdo a media altura", "=MAX($F$2:$F$%d)" % L, f_int),
            ("filader", "Fila del cruce derecho a media altura", "=MIN($G$2:$G$%d)" % L, f_int),
            ("tauizq", "τ del cruce izquierdo (fs)", f"=IF($M$8>0,INDEX($H$1:$H${L},$M$8),\"\")", f_calc),
            ("tauder", "τ del cruce derecho (fs)", f"=IF($M$9>0,INDEX($H$1:$H${L},$M$9),\"\")", f_calc),
            ("fwhm", "Ancho de la traza a media altura Δτ_AC (fs)",
             "=IF(AND($M$8>0,$M$9>0),ABS($M$11-$M$10),\"\")", f_key),
            ("taup", "Ancho de pulso τp = factor × Δτ_AC (fs)", "=IF(ISNUMBER($M$12),factor_forma*$M$12,\"\")", f_key),
            ("taug", "Referencia: τp si fuera gaussiano (fs)", "=IF(ISNUMBER($M$12),0.7071*$M$12,\"\")", f_calc),
            ("centro", "Centro de la traza (fs)",
             "=IF(ISNUMBER($M$12),($M$10+$M$11)/2,$M$7)", f_calc),
            ("ptsfw", "Lecturas dentro del ancho a media altura", f"=COUNTIF({rngE},\">=0.5\")", f_int),
            ("rizq", "Recorrido a la izquierda del centro (fs)", f"=IFERROR($M$15-MIN({rngC}),\"\")", f_calc),
            ("rder", "Recorrido a la derecha del centro (fs)", f"=IFERROR(MAX({rngC})-$M$15,\"\")", f_calc),
            ("monot", "¿Posiciones en un solo sentido?",
             f"=IF(COUNTIF($K$2:$K${L},0)>0,\"No: posición repetida\",IF(AND(COUNTIF($K$2:$K${L},1)>0,"
             f"COUNTIF($K$2:$K${L},-1)>0),\"No\",\"Sí\"))", f_lab),
            ("estado", "Estado de la traza",
             "=IF($M$2=0,\"Sin datos\",IF($M$2<>$M$3,\"Revisar: columnas A y B con distinto número de datos\","
             "IF($M$2<2*n_fondo+5,\"Revisar: muy pocos puntos\",IF($M$19<>\"Sí\","
             "\"Revisar: las posiciones no van en un solo sentido (ordenar A:B por la columna A)\",IF(OR($M$8=0,$M$9=0),"
             "\"Revisar: no se hallaron los dos cruces de media altura\",IF($M$16<pts_min,"
             "\"Aviso: pocas lecturas dentro del ancho; afinar el paso cerca del pico\",IF(MIN($M$17,$M$18)<rec_min*$M$12,"
             "\"Aviso: el barrido no llega a las alas\",\"OK\")))))))", f_keyt),
        ]
        for i, (k, et, fo, fm) in enumerate(panel):
            fila = i + 1
            P[k] = fila
            w.write(fila, 11, et, f_lab)
            w.write_formula(fila, 12, fo, fm)
        w.write(0, 11, f"Traza {nombre}", f_h)
        w.write_formula(0, 12, f"=\"Corriente \"&IFERROR(INDEX({q(S_REG)}!$C$5:$C${4 + N_TRAZAS},{num}),\"\")"
                        if num else "=\"Ejemplo\"", f_h)
        est = rc(P["estado"], 12)
        w.conditional_format(est, {"type": "cell", "criteria": "==", "value": '"OK"', "format": f_ok})
        w.conditional_format(est, {"type": "text", "criteria": "begins with", "value": "Aviso", "format": f_warn})
        w.conditional_format(est, {"type": "text", "criteria": "begins with", "value": "Revisar", "format": f_bad})
        # línea de media altura para la gráfica
        fl = len(panel) + 2
        w.write(fl, 11, "Línea de media altura (gráfica)", f_sub)
        w.write_formula(fl + 1, 11, "=IF(ISNUMBER($M$12),$M$10-$M$15,NA())", f_gray)
        w.write_formula(fl + 1, 12, "=IF(ISNUMBER($M$12),0.5,NA())", f_gray)
        w.write_formula(fl + 2, 11, "=IF(ISNUMBER($M$12),$M$11-$M$15,NA())", f_gray)
        w.write_formula(fl + 2, 12, "=IF(ISNUMBER($M$12),0.5,NA())", f_gray)
        w.write(fl + 4, 11, "Columnas A y B: desde la fila 2, sin filas vacías, en orden de posición (girando el "
                "micrómetro siempre en el mismo sentido). Posición en µm; potencia en W.", f_note)

        # columnas calculadas
        for i in range(max_filas):
            R = i + 2                                    # fila Excel
            ri = i                                       # índice base 0
            a, b = f"A{R}", f"B{R}"
            w.write_formula(ri + 1, 2, f"=IF(ISNUMBER({a}),{a}*conv,\"\")", f_gray)
            w.write_formula(ri + 1, 3, f"=IF(ISNUMBER({b}),AVERAGE(INDEX({rngB},MAX(1,{i + 1}-semi)):"
                                       f"INDEX({rngB},MIN($M$2,{i + 1}+semi))),\"\")", f_gray3)
            w.write_formula(ri + 1, 4, f"=IF(AND(ISNUMBER(D{R}),ISNUMBER($M$4)),(D{R}-$M$4)/($M$5-$M$4),\"\")", f_gray3)
            nxt = f"E{R + 1}"
            w.write_formula(ri + 1, 5, f"=IF(AND(ISNUMBER(E{R}),ISNUMBER({nxt}),ISNUMBER($M$6)),"
                                       f"IF(AND({R}<$M$6,E{R}<0.5,{nxt}>=0.5),{R},\"\"),\"\")", f_gray)
            w.write_formula(ri + 1, 6, f"=IF(AND(ISNUMBER(E{R}),ISNUMBER({nxt}),ISNUMBER($M$6)),"
                                       f"IF(AND({R}>=$M$6,E{R}>=0.5,{nxt}<0.5),{R},\"\"),\"\")", f_gray)
            w.write_formula(ri + 1, 7, f"=IF(OR(ISNUMBER(F{R}),ISNUMBER(G{R})),"
                                       f"C{R}+(0.5-E{R})*(C{R + 1}-C{R})/({nxt}-E{R}),\"\")", f_gray)
            w.write_formula(ri + 1, 8, f"=IF(AND(ISNUMBER(E{R}),ISNUMBER($M$15)),C{R}-$M$15,NA())", f_gray)
            w.write_formula(ri + 1, 9, f"=IF(ISNUMBER(E{R}),E{R},NA())", f_gray3)
            w.write_formula(ri + 1, 10, f"=IF(AND(ISNUMBER(C{R}),ISNUMBER(C{R + 1})),SIGN(C{R + 1}-C{R}),\"\")", f_gray)
        w.conditional_format(1, 8, max_filas, 9, {"type": "errors", "format": F(font_color="#D9D9D9")})
        w.conditional_format(fl + 1, 11, fl + 2, 12, {"type": "errors", "format": F(font_color="#D9D9D9")})
        if datos is not None:
            x, y = datos
            for i, xi in enumerate(x[:max_filas]):
                w.write_number(i + 1, 0, float(round(xi, 2)), f_in)
            for i, yi in enumerate(y[:max_filas]):
                w.write_number(i + 1, 1, float(yi), f_in)

        # gráfica cruda
        ch1 = wb.add_chart({"type": "scatter", "subtype": "straight_with_markers"})
        ch1.add_series({"name": "Datos crudos", "categories": [w.name, 1, 0, max_filas, 0],
                        "values": [w.name, 1, 1, max_filas, 1],
                        "marker": {"type": "circle", "size": 5, "fill": {"color": "#7F7F7F"},
                                   "border": {"color": "#7F7F7F"}}, "line": {"color": "#BFBFBF", "width": 0.75}})
        ch1.set_title({"name": "Traza cruda: potencia óptica contra desplazamiento", "name_font": {"size": 11}})
        ch1.set_x_axis({"name": "Posición del micrómetro (µm)", "major_gridlines": {"visible": False}})
        ch1.set_y_axis({"name": "Potencia óptica (W)", "major_gridlines": {"visible": True, "line": {"color": "#EEEEEE"}}})
        ch1.set_legend({"none": True})
        ch1.show_na_as_empty_cell()
        ch1.set_size({"width": 620, "height": 300})
        w.insert_chart(fl + 7, 11, ch1)
        # gráfica normalizada
        ch2 = wb.add_chart({"type": "scatter", "subtype": "straight_with_markers"})
        ch2.add_series({"name": "Traza normalizada", "categories": [w.name, 1, 8, max_filas, 8],
                        "values": [w.name, 1, 9, max_filas, 9],
                        "marker": {"type": "circle", "size": 5, "fill": {"color": "#1F3864"},
                                   "border": {"color": "#1F3864"}}, "line": {"color": "#8EA9DB", "width": 0.75}})
        ch2.add_series({"name": "Ancho a media altura", "categories": [w.name, fl + 1, 11, fl + 2, 11],
                        "values": [w.name, fl + 1, 12, fl + 2, 12],
                        "marker": {"type": "diamond", "size": 7, "fill": {"color": "#C00000"},
                                   "border": {"color": "#C00000"}}, "line": {"color": "#C00000", "width": 2}})
        ch2.set_title({"name": "Autocorrelación normalizada (sin fondo) contra retardo", "name_font": {"size": 11}})
        ch2.set_x_axis({"name": "Retardo respecto al centro (fs)", "major_gridlines": {"visible": False}})
        ch2.set_y_axis({"name": "Señal normalizada", "min": -0.1, "max": 1.1,
                        "major_gridlines": {"visible": True, "line": {"color": "#EEEEEE"}}})
        ch2.set_legend({"position": "bottom"})
        ch2.show_na_as_empty_cell()
        ch2.set_size({"width": 620, "height": 320})
        w.insert_chart(fl + 23, 11, ch2)
        w.print_area(0, 11, fl + 40, 17)               # panel y gráficas en una hoja
        w.fit_to_pages(1, 1)
        return P

    P = None
    for k, (w, n) in enumerate(zip(ws_tr, trazas)):
        d = datos_prueba[k] if datos_prueba else None
        P = hoja_traza(w, n, k + 1, d)
    hoja_traza(ws_ej, "Ejemplo", 0, traza_manual(54.0, semilla=7))
    ws_ej.merge_range(1, 13, 8, 16, "Traza manual simulada de un pulso sech² de 54 fs (7 A), con la guía de muestreo de "
                "Configuración: 16 lecturas cada 2 µm en la zona de solapamiento y 2 por ala, 20 en total; lecturas a 1 µm, "
                "ruido 1,5 %. Solo "
                "muestra cómo se ve una traza bien procesada; no entra al resumen.", f_note)
    ws_ej.set_tab_color("#A6A6A6")
    cel = lambda clave: f"$M${P[clave] + 1}"

    # ================================================================ Registro piloto
    w = ws_reg
    w.set_column("A:A", 6)
    w.set_column("B:C", 9)
    w.set_column("D:D", 11)
    w.set_column("E:F", 10)
    w.set_column("G:G", 13)
    w.set_column("H:H", 9)
    w.set_column("I:I", 34)
    w.set_column("J:K", 12)
    w.set_column("L:L", 44)
    w.write("A1", "Registro del piloto · sábado 24 de octubre de 2026", f_tit)
    w.write("A2", "Orden sorteado (semilla 20261024): 5 réplicas por corriente, 15 trazas. Medir en este orden; la "
            "traza N° k se pega en la hoja Rk.", f_sub)
    w.write("A3", "Amarillo = llenar. La duración se toma de la hora de inicio de una traza a la de la siguiente "
            "(incluye el cambio de corriente).", f_sub)
    w.write_row(3, 0, ["N°", "Hoja", "Corriente", "Corriente (A)", "Hora inicio", "Hora fin",
                       "Potencia media del láser (mW)", "¿Válida?", "Observaciones (causa si no es válida)",
                       "Δτ_AC (fs)", "τp (fs) · estimador elegido", "Estado de la traza", "Δ inicio a la siguiente (min)"], f_h)
    w.set_column("M:M", 12)
    w.set_row(3, 30)
    for k, cur in enumerate(orden_piloto()):
        R = 4 + k
        w.write_number(R, 0, k + 1, f_int)
        w.write_url(R, 1, f"internal:'{trazas[k]}'!A1", string=trazas[k], cell_format=f_linkt)
        w.write(R, 2, cur, f_in)
        w.write_formula(R, 3, f"=IF(C{R + 1}=\"\",\"\",IF(INDEX(corrientes,MATCH(C{R + 1},{{\"I1\",\"I2\",\"I3\"}},0))=\"\","
                              f"\"(falta)\",INDEX(corrientes,MATCH(C{R + 1},{{\"I1\",\"I2\",\"I3\"}},0))))", f_calc3)
        reg_k = (registro or {}).get(trazas[k], {})
        w.write(R, 4, reg_k.get("inicio", ""), f_in_t)
        w.write(R, 5, reg_k.get("fin", ""), f_in_t)
        w.write(R, 6, reg_k.get("potencia", ""), f_in_num)
        w.write(R, 7, reg_k.get("valida", "Sí"), f_in)
        w.write(R, 8, reg_k.get("obs", ""), f_in_txt)
        w.write_formula(R, 9, f"=IFERROR({q(trazas[k])}!{cel('fwhm')}+0,\"\")", f_link)
        aj = k + 5                                             # fila Excel de esta traza en Ajustes
        w.write_formula(R, 10, f"=IFERROR(IF(estimador=\"{ESTIMADORES[0]}\",{q(S_AJ)}!$I${aj}+0,"
                               f"IF(estimador=\"{ESTIMADORES[1]}\",{q(S_AJ)}!$J${aj}+0,"
                               f"IF(estimador=\"{ESTIMADORES[2]}\",{q(S_AJ)}!$K${aj}+0,{q(S_AJ)}!$N${aj}+0))),\"\")", f_link)
        if k < N_TRAZAS - 1:
            w.write_formula(R, 12, f"=IF(AND(H{R + 1}=\"Sí\",H{R + 2}=\"Sí\",ISNUMBER(E{R + 1}),ISNUMBER(E{R + 2})),"
                                   f"(E{R + 2}-E{R + 1})*1440,\"\")", f_gray)
        w.write_formula(R, 11, f"={q(trazas[k])}!{cel('estado')}", f_linkt)
    r1, r2 = 5, 4 + N_TRAZAS
    w.data_validation(f"C{r1}:C{r2}", {"validate": "list", "source": ["I1", "I2", "I3"]})
    w.data_validation(f"H{r1}:H{r2}", {"validate": "list", "source": ["Sí", "No"]})
    w.data_validation(f"E{r1}:F{r2}", {"validate": "time", "criteria": "between", "minimum": 0, "maximum": 0.99999,
                                         "error_message": "Escribir la hora como hh:mm, p. ej. 14:35"})
    rng_est = f"L{r1}:L{r2}"
    w.conditional_format(rng_est, {"type": "cell", "criteria": "==", "value": '"OK"', "format": f_ok})
    w.conditional_format(rng_est, {"type": "text", "criteria": "begins with", "value": "Aviso", "format": f_warn})
    w.conditional_format(rng_est, {"type": "text", "criteria": "begins with", "value": "Revisar", "format": f_bad})
    rr = r2 + 1
    w.write(rr, 8, "Duración media por traza (min)", f_labb)
    w.write_formula(rr, 9, f"=IF(COUNT(M{r1}:M{r2 - 1})=0,\"\",AVERAGE(M{r1}:M{r2 - 1}))", f_calc)
    w.write(rr, 10, "entre inicios de dos trazas válidas consecutivas (columna M)", f_note)
    if nota_registro:
        w.write(rr, 0, nota_registro, f_in_txt)
    w.write(rr + 1, 8, "Duración media de la toma sola (min)", f_lab)
    w.write_formula(rr + 1, 9, f"=IF(COUNT(F{r1}:F{r2})=0,\"\",SUMPRODUCT((F{r1}:F{r2}>0)*(E{r1}:E{r2}>0)*"
                               f"(F{r1}:F{r2}-E{r1}:E{r2}))*1440/SUMPRODUCT((F{r1}:F{r2}>0)*(E{r1}:E{r2}>0)))", f_calc)
    w.write(rr + 1, 10, "de inicio a fin de cada traza", f_note)
    wb.define_name("dur_piloto", f"={q(S_REG)}!$J${rr + 1}")
    w.set_row(rr + 3, 40)
    w.merge_range(rr + 3, 0, rr + 3, 11, "Reglas del protocolo (sección 5): una traza solo se marca «No» por una causa "
                  "anotada en el momento (el láser perdió los pulsos, la potencia salió del rango o el barrido se "
                  "interrumpió). Nunca por dar un valor inesperado. Toda traza descartada se repone con una nueva "
                  "(pegarla en la misma hoja y anotarlo).", f_note)
    w.freeze_panes(4, 0)

    # ================================================================ Resumen piloto
    w = ws_res
    w.set_column("A:A", 6)
    w.set_column("B:B", 10)
    w.set_column("C:E", 12)
    w.set_column("F:F", 2)
    w.set_column("G:N", 11)
    w.write("A1", "Resumen del piloto por corriente", f_tit)
    w.write("A2", "Solo cuentan las trazas con τp calculado y marcadas «Sí» en el Registro.", f_sub)
    w.write_row(3, 0, ["N°", "Corriente", "τp usable (fs)", "(τp − media)²", "Orden"], f_h)
    w.write_row(3, 6, ["x I1", "τp I1", "x I2", "τp I2", "x I3", "τp I3"], f_h)
    reg = q(S_REG)
    for k in range(N_TRAZAS):
        R = 4 + k
        E = R + 1
        w.write_formula(R, 0, f"={reg}!A{E}", f_int)
        w.write_formula(R, 1, f"={reg}!C{E}", f_linkt)
        w.write_formula(R, 2, f"=IF(AND(ISNUMBER({reg}!K{E}),{reg}!H{E}<>\"No\"),{reg}!K{E},\"\")", f_calc)
        w.write_formula(R, 3, f"=IF(ISNUMBER(C{E}),(C{E}-INDEX($C$22:$C$24,MATCH(B{E},$A$22:$A$24,0)))^2,0)", f_calc)
        w.write_formula(R, 4, f"=A{E}", f_int)
        for j, cur in enumerate(("I1", "I2", "I3")):
            w.write_formula(R, 6 + 2 * j, f"=IF(AND(B{E}=\"{cur}\",ISNUMBER(C{E})),{j + 1},NA())", f_gray)
            w.write_formula(R, 7 + 2 * j, f"=IF(AND(B{E}=\"{cur}\",ISNUMBER(C{E})),C{E},NA())", f_gray)
    a1, a2 = 5, 4 + N_TRAZAS
    w.conditional_format(4, 6, a2 - 1, 11, {"type": "errors", "format": F(font_color="#D9D9D9")})
    w.merge_range(19, 6, 19, 11, "Columnas G:L: solo alimentan las gráficas (#N/A = punto que no se dibuja).", f_note)
    w.write_row(20, 0, ["Corriente", "n", "Media τp (fs)", "S (fs)", "CV (%)"], f_h)
    w.write_row(20, 6, ["Corriente (A)", "τ0 fábrica (fs)", "Media − τ0 (fs)", "x"], f_h)
    for j, cur in enumerate(("I1", "I2", "I3")):
        R = 21 + j
        E = R + 1
        w.write(R, 0, cur, f_labb)
        w.write_formula(R, 1, f"=COUNTIFS($B${a1}:$B${a2},A{E},$C${a1}:$C${a2},\">0\")", f_int)
        w.write_formula(R, 2, f"=IF(B{E}=0,\"\",AVERAGEIFS($C${a1}:$C${a2},$B${a1}:$B${a2},A{E}))", f_key)
        w.write_formula(R, 3, f"=IF(B{E}<2,\"\",SQRT(SUMIFS($D${a1}:$D${a2},$B${a1}:$B${a2},A{E})/(B{E}-1)))", f_key)
        w.write_formula(R, 4, f"=IF(ISNUMBER(D{E}),100*D{E}/C{E},\"\")", f_calc)
        w.write_formula(R, 6, f"=IF(I{j + 1}_A=\"\",\"\",I{j + 1}_A)", f_link)
        w.write_formula(R, 7, f"=IF(ISNUMBER(tau0_{j + 1}),tau0_{j + 1},NA())", f_link)
        w.write_formula(R, 8, f"=IF(AND(ISNUMBER(C{E}),ISNUMBER(H{E})),C{E}-H{E},\"\")", f_calc)
        w.write_number(R, 9, j + 1, f_gray)
        wb.define_name(f"n_{cur}", f"={q(S_RES)}!$B${E}")
        wb.define_name(f"S_{cur}", f"={q(S_RES)}!$D${E}")
    # gráfica de puntos por corriente
    ch = wb.add_chart({"type": "scatter"})
    colores = ("#1F3864", "#C55A11", "#548235")
    for j, cur in enumerate(("I1", "I2", "I3")):
        ch.add_series({"name": cur, "categories": [S_RES, 4, 6 + 2 * j, a2 - 1, 6 + 2 * j],
                       "values": [S_RES, 4, 7 + 2 * j, a2 - 1, 7 + 2 * j],
                       "marker": {"type": "circle", "size": 7, "fill": {"color": colores[j]},
                                  "border": {"color": colores[j]}}})
    ch.add_series({"name": "τ0 fábrica", "categories": [S_RES, 21, 9, 23, 9], "values": [S_RES, 21, 7, 23, 7],
                   "marker": {"type": "dash", "size": 14, "fill": {"color": "#C00000"}, "border": {"color": "#C00000"}}})
    ch.set_title({"name": "Diagrama de puntos: τp por corriente y valor de fábrica", "name_font": {"size": 11}})
    ch.set_x_axis({"name": "Corriente (1 = I1, 2 = I2, 3 = I3)", "min": 0, "max": 4, "major_unit": 1,
                   "major_gridlines": {"visible": False}})
    ch.set_y_axis({"name": "τp (fs)", "major_gridlines": {"visible": True, "line": {"color": "#EEEEEE"}}})
    ch.set_legend({"position": "bottom"})
    ch.show_na_as_empty_cell()
    ch.set_size({"width": 520, "height": 300})
    w.insert_chart("A28", ch)
    # τp contra orden
    ch = wb.add_chart({"type": "scatter"})
    for j, cur in enumerate(("I1", "I2", "I3")):
        ch.add_series({"name": cur, "categories": [S_RES, 4, 4, a2 - 1, 4],
                       "values": [S_RES, 4, 7 + 2 * j, a2 - 1, 7 + 2 * j],
                       "marker": {"type": "circle", "size": 7, "fill": {"color": colores[j]},
                                  "border": {"color": colores[j]}}})
    ch.set_title({"name": "τp contra el orden de medición (independencia)", "name_font": {"size": 11}})
    ch.set_x_axis({"name": "Orden de medición", "min": 0, "max": N_TRAZAS + 1, "major_unit": 1,
                   "major_gridlines": {"visible": False}})
    ch.set_y_axis({"name": "τp (fs)", "major_gridlines": {"visible": True, "line": {"color": "#EEEEEE"}}})
    ch.set_legend({"position": "bottom"})
    ch.show_na_as_empty_cell()
    ch.set_size({"width": 520, "height": 300})
    w.insert_chart("H28", ch)
    w.print_area("A1:P45")
    w.set_landscape()
    w.fit_to_pages(1, 1)
    w.merge_range("A25:N26", "Con solo 5 datos por corriente no se prueba normalidad: el piloto solo estima S. La "
                  "normalidad se revisa con los datos de la semana (protocolo, sección 6).", f_note)

    # ================================================================ Réplicas
    w = ws_rep
    w.set_column("A:A", 52)
    w.set_column("B:E", 13)
    w.set_column("F:F", 2)
    w.set_column("G:G", 60)
    w.write("A1", "Número de réplicas de la semana (protocolo, sección 4)", f_tit)
    w.write("A2", "Todo es automático. Celdas verdes = resultados clave. Montgomery, sec. 2-4.2 y curvas OC (apéndice V).",
            f_sub)
    # --- paso 1: sigma
    w.write_row(3, 0, ["Paso 1 · σ conservadora", "I1", "I2", "I3", "Usado"], f_h)
    w.write(4, 0, "n del piloto (trazas válidas)", f_lab)
    w.write(5, 0, "S del piloto (fs)", f_lab)
    for j, cur in enumerate(("I1", "I2", "I3")):
        w.write_formula(4, 1 + j, f"=n_{cur}", f_link)
        w.write_formula(5, 1 + j, f"=IF(ISNUMBER(S_{cur}),S_{cur},\"\")", f_link)
    w.write_formula(5, 4, "=IF(COUNT(B6:D6)=0,\"\",MAX(B6:D6))", f_calc)
    w.write(6, 0, "Corriente con la S mayor", f_lab)
    w.write_formula(6, 4, "=IF(ISNUMBER(E6),INDEX($B$4:$D$4,MATCH(E6,B6:D6,0)),\"\")", f_linkt)
    w.write(7, 0, "n con que se estimó esa S", f_lab)
    w.write_formula(7, 4, "=IF(ISNUMBER(E6),INDEX(B5:D5,MATCH(E6,B6:D6,0)),\"\")", f_int)
    w.write(8, 0, "Factor del límite superior: √[(n − 1)/χ²(1 − conf; n − 1)]", f_lab)
    w.write_formula(8, 4, "=IF(AND(ISNUMBER(E8),E8>=2),SQRT((E8-1)/CHIINV(conf_sigma,E8-1)),\"\")", f_calc3)
    w.write(9, 0, "σ para el diseño = factor × S mayor (fs)", f_labb)
    w.write_formula(9, 4, "=IF(ISNUMBER(E9),E9*E6,\"\")", f_key)
    w.write(8, 6, "Con n = 5 y conf = 0,80 el factor es 1,557 (el «1,56» del protocolo).", f_note)
    # --- paso 2: delta
    w.write_row(11, 0, ["Paso 2 · δ que importa detectar", "I1", "I2", "I3", "Usado"], f_h)
    w.write(12, 0, "Regla que aplica (primer dato del fabricante que exista)", f_lab)
    w.write(13, 0, "δ según la regla (fs)", f_lab)
    w.write(14, 0, "δ con el piso 2u (fs)", f_labb)
    for k in (1, 2, 3):
        col = k
        regla = (f"=IF(ISNUMBER(tol_{k}),\"1 · tolerancia\",IF(AND(ISNUMBER(tau0_{k}),ISNUMBER(taumax_{k}),taumax_{k}>tau0_{k}),"
                 f"\"2 · τmáx − τ0\",IF(ISNUMBER(banda_{k}),\"3 · banda\",IF(AND(ISNUMBER(estab_{k}),ISNUMBER(tau0_{k})),"
                 f"\"4 · estabilidad\",\"5 · por defecto\"))))")
        dval = (f"=IF(ISNUMBER(tol_{k}),tol_{k},IF(AND(ISNUMBER(tau0_{k}),ISNUMBER(taumax_{k}),taumax_{k}>tau0_{k}),taumax_{k}-tau0_{k},"
                f"IF(ISNUMBER(banda_{k}),banda_{k},IF(AND(ISNUMBER(estab_{k}),ISNUMBER(tau0_{k})),"
                f"3*estab_{k}/100*tau0_{k},delta_def))))")
        w.write_formula(12, col, regla, F(border=1, font_size=8, text_wrap=True))
        w.write_formula(13, col, dval, f_calc)
        w.write_formula(14, col, f"=IF(ISNUMBER(incert_{k}),MAX({rc(13, col)},2*incert_{k}),{rc(13, col)})", f_calc)
    w.write_formula(14, 4, "=MIN(B15:D15)", f_key)
    w.write(14, 6, "Si δ cambia por corriente, se usa la menor (la que exige más réplicas). La regla 2 necesita τ0 "
            "medido; sin él cae a la 5 (10 fs).", f_note)
    w.set_row(12, 26)
    # --- paso 3: n
    w.write_row(16, 0, ["Paso 3 · Número de réplicas por corriente", "", "", "", "Valor"], f_h)
    w.write(17, 0, "α · potencia · colas", f_lab)
    w.write_formula(17, 4, "=alfa&\" · \"&potencia&\" · \"&colas", f_linkt)
    w.write(18, 0, "d = δ/σ", f_lab)
    w.write_formula(18, 4, "=IF(ISNUMBER(E10),E15/E10,\"\")", f_calc3)
    w.write(19, 0, "n necesario por corriente (curva OC, t de una muestra)", f_labb)
    w.write_formula(19, 4, "=IF(ISNUMBER(E19),IF(COUNT($N$5:$N$103)=0,\"> 100\",MIN($N$5:$N$103)),\"\")", f_keyi)
    w.write(20, 0, "Potencia que se logra con ese n", f_lab)
    w.write_formula(20, 4, "=IF(ISNUMBER(E20),INDEX($M$5:$M$103,E20-1),\"\")", f_calc3)
    w.write(19, 6, "Verificar en Minitab: Estadísticas › Potencia y tamaño de la muestra › t de 1 muestra, con "
            "Diferencia = δ, Desv. Est. = σ y Potencia = 0,90.", f_note)
    # --- paso 4: calendario
    w.write_row(22, 0, ["Paso 4 · Reparto en la semana", "3 días", "4 días", "", ""], f_h)
    w.write(23, 0, "Minutos por traza (incluye cambio de corriente)", f_lab)
    w.write_formula(23, 1, "=IF(ISNUMBER(min_traza_man),min_traza_man,dur_piloto)", f_link)
    w.write(24, 0, "Trazas que caben por día", f_lab)
    w.write_formula(24, 1, "=IF(AND(ISNUMBER(B24),B24>0),INT(horas_dia*60/B24),\"\")", f_int)
    w.write(25, 0, "Réplicas por corriente y por día = ⌈n / días⌉", f_labb)
    w.write(26, 0, "Trazas por día = 3 × réplicas por día", f_lab)
    w.write(27, 0, "¿Caben en el día?", f_lab)
    w.write(28, 0, "n por corriente del plan (réplicas por día × días)", f_labb)
    w.write(29, 0, "Si no caben: réplicas por corriente y día que sí caben", f_lab)
    w.write(30, 0, "n por corriente que sí cabe", f_lab)
    w.write(31, 0, "n por corriente que se mide", f_labb)
    w.write(32, 0, "δ detectable δdet con ese n (fs)", f_labb)
    w.write(33, 0, "¿δdet ≤ δ?", f_lab)
    w.write(34, 0, "Total de trazas de la semana", f_lab)
    for j, dias in enumerate((3, 4)):
        c = 1 + j
        col = "BC"[j]
        w.write_formula(24, c, "=B25", f_int) if j else None
        w.write_formula(25, c, f"=IF(ISNUMBER($E$20),ROUNDUP($E$20/{dias},0),\"\")", f_keyi)
        w.write_formula(26, c, f"=IF(ISNUMBER({col}26),3*{col}26,\"\")", f_int)
        w.write_formula(27, c, f"=IF(AND(ISNUMBER({col}27),ISNUMBER($B$25)),IF({col}27<=$B$25,\"Sí\",\"No\"),"
                               f"\"(falta cronometraje)\")", f_linkt)
        w.write_formula(28, c, f"=IF(ISNUMBER({col}26),{col}26*{dias},\"\")", f_int)
        w.write_formula(29, c, f"=IF({col}28=\"No\",INT($B$25/3),\"\")", f_int)
        w.write_formula(30, c, f"=IF(ISNUMBER({col}30),{col}30*{dias},\"\")", f_int)
        w.write_formula(31, c, f"=IF(ISNUMBER({col}31),{col}31,{col}29)", f_keyi)
        w.write_formula(32, c, f"=IF(AND(ISNUMBER({col}32),{col}32>=2,ISNUMBER($E$10)),"
                               f"INDEX($S$5:$S$103,{col}32-1)*$E$10,\"\")", f_calc)
        w.write_formula(33, c, f"=IF(ISNUMBER({col}33),IF({col}33<=$E$15+0.005,\"Sí\",\"No: declararlo en el informe\"),\"\")",
                        f_linkt)
        w.write_formula(34, c, f"=IF(ISNUMBER({col}32),3*{col}32,\"\")", f_int)
    w.write(25, 6, "Se redondea hacia arriba para repartir igual entre días; el n del plan puede quedar un poco por "
            "encima del necesario.", f_note)
    w.write(32, 6, "Mínima diferencia que el plan detecta con α y potencia fijados (la «salida» de la curva OC con n "
            "fijo).", f_note)
    w.write(36, 0, "Decisión: días de medición (3 = lun–mié; 4 = usar también el jueves 29)", f_labb)
    w.write(36, 1, 3, F(border=2, bg_color="#FFF2CC", font_color="#0000FF", num_format="0", bold=True))
    w.data_validation(36, 1, 36, 1, {"validate": "list", "source": [3, 4]})
    wb.define_name("dias", f"={q(S_REP)}!$B$37")
    w.write(37, 0, "Réplicas por corriente y por día del plan", f_labb)
    w.write_formula(37, 1, "=IF(dias=4,IF(ISNUMBER(C32),C32/4,\"\"),IF(ISNUMBER(B32),B32/3,\"\"))", f_keyi)
    wb.define_name("rpd", f"={q(S_REP)}!$B$38")
    w.write(38, 0, "n por corriente del plan", f_labb)
    w.write_formula(38, 1, "=IF(ISNUMBER(B38),B38*dias,\"\")", f_keyi)
    w.print_area("A1:G40")
    w.fit_to_pages(1, 0)
    w.write(37, 6, f"La hoja Plan semana arma el orden sorteado de cada día con este número (máximo {MAX_RPD}).", f_note)

    # --- tabla de potencia (curva OC numérica)
    w.set_column("H:H", 2)
    w.set_column("I:S", 9)
    w.write("I3", "Curva OC numérica: potencia de la t de una muestra (aproximación normal de la t no central)", F(bold=True))
    w.write_row(3, 8, ["n", "ν", "t crítico", "λ = d√n", "Potencia", "n si ≥ 1−β",
                       "d₀", "d₁", "d₂", "d₃", "d_det"], f_h)
    # columnas: I n, J nu, K tcrit, L lam, M pot, N flag, O..R iteraciones Newton, S d_det
    w.set_column("O:R", 9, None, {"level": 1, "hidden": True})

    def pot(d, nn, nu, tc):
        """Fórmula (texto) de la potencia aproximada para d (texto)."""
        s = f"SQRT(1+{tc}^2/(2*{nu}))"
        cc = f"(1-1/(4*{nu}))"
        lam = f"({d})*SQRT({nn})"
        p = f"(1-NORMSDIST(({tc}*{cc}-{lam})/{s}))"
        return f"({p}+IF(colas=2,NORMSDIST((-{tc}*{cc}-{lam})/{s}),0))"

    for i, nn in enumerate(range(2, 101)):
        R = 4 + i
        E = R + 1
        w.write_number(R, 8, nn, f_int)
        w.write_formula(R, 9, f"=I{E}-1", f_int)
        w.write_formula(R, 10, f"=IF(colas=2,TINV(alfa,J{E}),TINV(2*alfa,J{E}))", f_calc3)
        w.write_formula(R, 11, f"=IF(ISNUMBER($E$19),$E$19*SQRT(I{E}),\"\")", f_calc3)
        w.write_formula(R, 12, f"=IF(ISNUMBER($E$19),{pot('$E$19', f'I{E}', f'J{E}', f'K{E}')},\"\")", f_calc3)
        w.write_formula(R, 13, f"=IF(ISNUMBER(M{E}),IF(M{E}>=potencia,I{E},\"\"),\"\")", f_int)
        # δ detectable: Newton sobre la potencia aproximada, partiendo de (t + tβ)/√n
        w.write_formula(R, 14, f"=(K{E}+TINV(2*(1-potencia),J{E}))/SQRT(I{E})", f_calc3)
        for it, (src, dst) in enumerate((("O", 15), ("P", 16), ("Q", 17))):
            dk = f"{src}{E}"
            h = f"(0.001*{dk})"
            f0 = pot(dk, f"I{E}", f"J{E}", f"K{E}")
            fp = pot(f"{dk}+{h}", f"I{E}", f"J{E}", f"K{E}")
            fm = pot(f"{dk}-{h}", f"I{E}", f"J{E}", f"K{E}")
            w.write_formula(R, dst, f"=IFERROR({dk}-({f0}-potencia)/(({fp}-{fm})/(2*{h})),{dk})", f_calc3)
        w.write_formula(R, 18, f"=R{E}", f_calc3)

    # ================================================================ Ajustes
    w = ws_aj
    w.set_column("A:A", 5)
    w.set_column("B:D", 9)
    w.set_column("E:P", 11)
    w.set_column("R:U", 10, None, {"level": 1, "hidden": True})
    w.write("A1", "Cuatro estimadores de τp por traza: ancho a media altura (FWHM) y ajuste por mínimos cuadrados, "
            "para pulso sech² y gaussiano", f_tit)
    w.write("A2", "Columnas E–J: fórmulas sobre las hojas R01–R15. Columnas K–P: ajuste no lineal de B + A·g(τ − τc; τp) "
            "a todos los puntos, calculado con scipy (scripts/trabajo-practico/analiza_piloto.py); son valores, no "
            "fórmulas. El Resumen usa el estimador elegido en Configuración.", f_sub)
    w.write_row(3, 0, ["N°", "Hoja", "Corriente", "¿Válida?", "Puntos", "Paso (fs)", "Lecturas en el FWHM", "Δτ_AC (fs)",
                       "τp FWHM sech² (fs)", "τp FWHM gauss (fs)", "τp ajuste sech² (fs)", "± SE", "R² sech²",
                       "τp ajuste gauss (fs)", "± SE", "R² gauss"], f_h)
    w.set_row(3, 42)
    f_val = F(border=1, bg_color="#EDEDED", num_format="0.00")
    f_val3 = F(border=1, bg_color="#EDEDED", num_format="0.000")
    for k, hoja in enumerate(trazas):
        R = 4 + k
        E = R + 1
        sh = q(hoja)
        w.write_formula(R, 0, f"={reg}!A{E}", f_int)
        w.write_url(R, 1, f"internal:'{hoja}'!A1", string=hoja, cell_format=f_linkt)
        w.write_formula(R, 2, f"={reg}!C{E}", f_linkt)
        w.write_formula(R, 3, f"={reg}!H{E}", f_linkt)
        w.write_formula(R, 4, f"={sh}!$M$2", f_int)
        w.write_formula(R, 5, f"=IF({sh}!$M$2>1,(MAX({sh}!$C$2:$C${L})-MIN({sh}!$C$2:$C${L}))/({sh}!$M$2-1),\"\")", f_calc)
        w.write_formula(R, 6, f"=IF({sh}!$M$2>0,{sh}!{cel('ptsfw')},\"\")", f_int)
        w.write_formula(R, 7, f"=IFERROR({sh}!{cel('fwhm')}+0,\"\")", f_link)
        w.write_formula(R, 8, f"=IF(ISNUMBER(H{E}),factor_forma*H{E},\"\")", f_calc)
        w.write_formula(R, 9, f"=IF(ISNUMBER(H{E}),0.7071*H{E},\"\")", f_calc)
        a = (ajustes or {}).get(hoja)
        if a:
            w.write_number(R, 10, a["fit_sech"]["tau_p"], f_val)
            w.write_number(R, 11, a["fit_sech"]["se_tau_p"], f_val)
            w.write_number(R, 12, a["fit_sech"]["r2"], f_val3)
            w.write_number(R, 13, a["fit_gauss"]["tau_p"], f_val)
            w.write_number(R, 14, a["fit_gauss"]["se_tau_p"], f_val)
            w.write_number(R, 15, a["fit_gauss"]["r2"], f_val3)
        else:
            for c in range(10, 16):
                w.write_blank(R, c, None, f_val)
        # helpers: (x − media de su corriente)² por estimador, solo válidas
        for j, col in enumerate("IKJN"):
            hc = 17 + j
            w.write_formula(R, hc, f"=IF(AND($D{E}=\"Sí\",ISNUMBER({col}{E})),({col}{E}-AVERAGEIFS(${col}$5:${col}$19,"
                                   f"$C$5:$C$19,$C{E},$D$5:$D$19,\"Sí\"))^2,0)", f_gray)
    w.write_row(3, 17, ["h FWHM sech²", "h ajuste sech²", "h FWHM gauss", "h ajuste gauss"], f_h)
    w.write("A21", "Resumen por corriente y estimador (solo trazas válidas)", f_labb)
    w.write_row(21, 0, ["Corriente", "τ0 fábrica", "n", "FWHM sech²: media", "S", "Ajuste sech²: media", "S",
                        "FWHM gauss: media", "S", "Ajuste gauss: media", "S"], f_h)
    w.set_row(21, 30)
    for j, cur in enumerate(("I1", "I2", "I3")):
        R = 22 + j
        E = R + 1
        w.write(R, 0, cur, f_labb)
        w.write_formula(R, 1, f"=tau0_{j + 1}", f_link)
        w.write_formula(R, 2, f"=COUNTIFS($C$5:$C$19,A{E},$D$5:$D$19,\"Sí\",$I$5:$I$19,\">0\")", f_int)
        for m, (col, hcol) in enumerate((("I", "R"), ("K", "S"), ("J", "T"), ("N", "U"))):
            w.write_formula(R, 3 + 2 * m, f"=IF(C{E}=0,\"\",AVERAGEIFS(${col}$5:${col}$19,$C$5:$C$19,A{E},$D$5:$D$19,\"Sí\"))", f_key)
            w.write_formula(R, 4 + 2 * m, f"=IF(C{E}<2,\"\",SQRT(SUMIFS(${hcol}$5:${hcol}$19,$C$5:$C$19,A{E},$D$5:$D$19,\"Sí\")/(C{E}-1)))", f_calc)
    w.write("A27", "El ajuste gaussiano y el sech² dan anchos distintos para la misma traza porque suponen formas de pulso "
            "distintas; la hoja de pruebas de Calmar usa sech² (factor 0,65). Se reportan ambos porque la hoja del "
            "equipo no declara la forma.", f_note)
    w.freeze_panes(4, 2)

    # ================================================================ Plan semana
    w = ws_plan
    w.set_column("A:A", 8)
    w.write("A1", "Plan de medición de la semana · orden sorteado por día", f_tit)
    w.write_formula("A2", "=\"Réplicas por corriente y por día: \"&IF(ISNUMBER(rpd),rpd,\"(se calcula en la hoja Réplicas)\")"
                          "&\" · días: \"&dias&\" · semilla: \"&semilla", f_sub)
    w.merge_range("A3:I3", "Al fijar el plan el sábado por la noche, copiar estas columnas y pegarlas como valores en "
                  "las hojas de registro de cada día: si se cambia un parámetro después, el orden cambia.", f_note)
    w.set_row(2, 28)
    nombres_dia = ["Lun 26 oct", "Mar 27 oct", "Mié 28 oct", "Jue 29 oct"]
    nmax = 3 * MAX_RPD
    # bloques por día: Posición | Corriente | Corriente (A); helpers ocultos a la derecha
    w.write(4, 0, "Posición", f_h)
    for i in range(nmax):
        w.write_formula(5 + i, 0, f"=IF(ISNUMBER(rpd),IF({i + 1}<=3*rpd,{i + 1},\"\"),\"\")", f_int)
    hc = 12                                                    # columna inicial de auxiliares
    for dd in range(DIAS_MAX):
        c = 1 + 2 * dd
        w.set_column(c, c, 12)
        w.set_column(c + 1, c + 1, 12)
        w.merge_range(3, c, 3, c + 1, nombres_dia[dd], f_h)
        w.write(4, c, "Corriente", f_h)
        w.write(4, c + 1, "Valor (A)", f_h)
        hk, hl, hr = hc + 3 * dd, hc + 3 * dd + 1, hc + 3 * dd + 2      # clave, etiqueta, rango
        w.write(4, hk, f"clave d{dd + 1}", f_h)
        w.write(4, hl, f"etiq d{dd + 1}", f_h)
        w.write(4, hr, f"rango d{dd + 1}", f_h)
        colk = xlsxwriter.utility.xl_col_to_name(hk)
        coll = xlsxwriter.utility.xl_col_to_name(hl)
        colr = xlsxwriter.utility.xl_col_to_name(hr)
        r1, r2 = 6, 5 + nmax
        for i in range(nmax):
            E = 6 + i
            k = i + 1
            w.write_formula(5 + i, hk, f"=MOD(ABS(SIN(semilla*12.9898+{k}*78.233+{dd + 1}*37.719))*43758.5453,1)", f_gray3)
            w.write_formula(5 + i, hl, f"=\"I\"&(MOD({k}-1,3)+1)", f_gray)
            w.write_formula(5 + i, hr, f"=IF(ISNUMBER($A{E}),IF({dd + 1}<=dias,COUNTIFS(${colk}${r1}:${colk}${r2},"
                                       f"\"<\"&{colk}{E},$A${r1}:$A${r2},\">0\")+1,\"\"),\"\")", f_gray)
            w.write_formula(5 + i, c, f"=IF(ISNUMBER($A{E}),IF({dd + 1}<=dias,IFERROR(INDEX(${coll}${r1}:${coll}${r2},"
                                      f"MATCH($A{E},${colr}${r1}:${colr}${r2},0)),\"\"),\"—\"),\"\")", f_linkt)
            w.write_formula(5 + i, c + 1, f"=IF(OR({rc(5 + i, c)}=\"\",{rc(5 + i, c)}=\"—\"),\"\","
                                          f"IF(INDEX(corrientes,MATCH({rc(5 + i, c)},{{\"I1\",\"I2\",\"I3\"}},0))=\"\",\"\","
                                          f"INDEX(corrientes,MATCH({rc(5 + i, c)},{{\"I1\",\"I2\",\"I3\"}},0))))", f_calc3)
        w.conditional_format(5, c, 4 + nmax, c, {"type": "cell", "criteria": "==", "value": '"I1"',
                                                  "format": F(bg_color="#D9E1F2", border=1)})
        w.conditional_format(5, c, 4 + nmax, c, {"type": "cell", "criteria": "==", "value": '"I2"',
                                                  "format": F(bg_color="#FBE5D6", border=1)})
        w.conditional_format(5, c, 4 + nmax, c, {"type": "cell", "criteria": "==", "value": '"I3"',
                                                  "format": F(bg_color="#E2F0D9", border=1)})
    w.set_column(hc, hc + 3 * DIAS_MAX - 1, 9, None, {"level": 1, "hidden": True})
    w.write(4, 10, "Columnas auxiliares del sorteo: grupo M:X, oculto", f_sub)
    w.freeze_panes(5, 1)

    # ================================================================ Instrucciones
    w = ws_ins
    w.set_column("A:A", 4)
    w.set_column("B:B", 110)
    w.write("B1", "Piloto del autocorrelador · cálculo del número de réplicas", f_tit)
    w.write("B2", "Trabajo práctico · Diseño de Experimentos Avanzados (3008475) · UNAL 2026-II · "
            "J. E. Rodríguez Ochoa y S. Zapata Henao · sábado 24 de octubre de 2026", f_sub)
    pasos = [
        ("Antes de medir", ""),
        ("1", "Configuración: la posición va en µm y la potencia óptica en W. Revisar que el haz recorra 2 veces el "
              "desplazamiento (retrorreflector: τ = 2Δx/c). Las corrientes ya están puestas según la hoja de pruebas de "
              "fábrica: I1 = 5 A (τ0 = 107 fs), I2 = 6 A (76 fs), I3 = 7 A (54 fs)."),
        ("2", "Registro piloto: el orden de las 15 trazas ya está sorteado (5 por corriente). Medir en ese orden."),
        ("Durante el piloto", ""),
        ("3", "En la hoja Rk (k = número de la traza en el Registro), escribir desde la fila 2 la posición del "
              "micrómetro (columna A, en µm) y la potencia óptica (columna B, en W), una lectura por fila, en orden de posición. "
              "Con el montaje actual son unas 20 lecturas (16–17 en la zona de solapamiento y 2 por ala); caben hasta %d. "
              "La guía de muestreo (Configuración, numeral 7) da el paso y la posición de las alas para cada corriente."
              % max_filas),
        ("4", "Mirar el Estado de la traza (celda M20) y las dos gráficas a la derecha. «OK» = la traza se procesó bien. "
              "Un «Aviso» no invalida la traza pero va al informe; un «Revisar» hay que mirarlo antes de seguir."),
        ("5", "En el Registro, anotar la hora de inicio y de fin (hh:mm), la potencia media del láser en mW y, si la "
              "traza se descarta por una causa anotada, «No» y la causa."),
        ("Al terminar (sábado por la noche)", ""),
        ("6", "Resumen piloto: media, S y CV de τp por corriente y las gráficas de puntos y de τp contra el orden."),
        ("7", "Configuración, numeral 5: llenar los datos del fabricante que se tengan (si no hay ninguno, δ = 10 fs). "
              "Revisar también las horas útiles por día."),
        ("8", "Réplicas: el n por corriente sale solo (σ = 1,56·S mayor, δ por la regla, α = 0,05, potencia 0,90). "
              "Elegir 3 o 4 días en la celda B37 y revisar si caben y el δ detectable."),
        ("9", "Plan semana: orden sorteado de cada día. Copiarlo como valores a las hojas de registro de la semana."),
        ("Cómo se calcula τp en cada traza", ""),
        ("·", "Retardo τ = (pasos × desplazamiento en metros / c) × 10¹⁵ fs; con 2 pasos, 6,67 fs por cada µm del micrómetro."),
        ("·", "Se resta el fondo (promedio de los primeros y últimos N puntos), se normaliza al máximo y se buscan los "
              "dos cruces a 0,5 más cercanos al pico, interpolando linealmente entre puntos."),
        ("·", "Δτ_AC = distancia entre los cruces; τp = 0,65 × Δτ_AC (factor de Calmar, sech²). Es el ancho a media "
              "altura medido directamente, sin ajuste por mínimos cuadrados. Con 20 lecturas según la guía (≈ 6 dentro "
              "del ancho) la simulación da un sesgo de +0,4 a +0,9 fs (igual para todas las trazas de una corriente), "
              "frente a una dispersión de 0,7 a 1,6 fs entre trazas."),
        ("·", "El número de réplicas usa la aproximación normal de la t no central; se comparó con el cálculo exacto "
              "(scipy) para d entre 0,3 y 4: da el mismo n en 147 de 148 casos (el otro difiere en 1). Confirmar el "
              "resultado final en Minitab."),
        ("Colores", ""),
        ("·", "Amarillo con letra azul = se escribe. Verde = resultado clave. Letra verde = viene de otra hoja. "
              "Gris = cálculo intermedio. La hoja Ejemplo tiene una traza manual simulada de 54 fs (7 A) para ver cómo debe quedar."),
    ]
    r = 3
    for a, b in pasos:
        if b == "":
            w.write(r, 1, a, F(bold=True, font_color="#1F3864", font_size=11, top=1))
        else:
            w.write(r, 0, a, F(bold=True, align="center", valign="top"))
            w.write(r, 1, b, f_txt)
        r += 1
    w.activate()

    wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--prueba", type=Path, help="genera una copia pequeña con datos sintéticos")
    ap.add_argument("--desde", type=Path, help="libro ya usado: copia sus datos, horas y validez")
    ap.add_argument("--analisis", type=Path, help="JSON de analiza_piloto.py con los ajustes")
    ap.add_argument("--salida", type=Path, help="ruta del libro generado con --desde")
    ap.add_argument("--descartar", default="", help="hojas a marcar como no válidas, p. ej. R02,R03")
    ap.add_argument("--excluir", default="", help="hojas a marcar como no válidas por no mostrar la autocorrelación")
    args = ap.parse_args()
    if args.desde:
        import json
        import openpyxl
        wb = openpyxl.load_workbook(args.desde)
        reg = wb["Registro piloto"]
        descartar = set(filter(None, args.descartar.split(",")))
        excluir = set(filter(None, args.excluir.split(",")))
        registro, datos = {}, []
        for k in range(N_TRAZAS):
            hoja = f"R{k + 1:02d}"
            fila = 5 + k
            registro[hoja] = {"inicio": reg.cell(fila, 5).value or "", "fin": reg.cell(fila, 6).value or "",
                              "potencia": reg.cell(fila, 7).value or "",
                              "valida": "No" if hoja in descartar | excluir else (reg.cell(fila, 8).value or "Sí"),
                              "obs": "Descartada por error de medición (decisión del equipo, 24 oct)" if hoja in descartar
                              else ("Excluida: sin señal de autocorrelación (potencia 10 veces menor que las demás trazas, sin pico)." if hoja in excluir
                                    else (reg.cell(fila, 9).value or ""))}
            ws = wb[hoja]
            x = [ws.cell(r, 1).value for r in range(2, 302) if ws.cell(r, 1).value is not None]
            y = [ws.cell(r, 2).value for r in range(2, 302) if ws.cell(r, 2).value is not None]
            datos.append((x, y))
        nota = reg["A21"].value
        aj = json.load(open(args.analisis))["trazas"] if args.analisis else None
        build(args.salida, 300, datos, registro=registro, ajustes=aj, nota_registro=nota)
        print(args.salida)
    elif args.prueba:
        rng = np.random.default_rng(3)
        medias = dict(zip(("I1", "I2", "I3"), map(float, TAU0_FAB)))
        datos = [traza_manual(medias[c] + 3 * rng.standard_normal(), semilla=k + 10)
                 for k, c in enumerate(orden_piloto())]
        build(args.prueba, 300, datos)
        print(args.prueba)
    else:
        SALIDA.parent.mkdir(parents=True, exist_ok=True)
        build(SALIDA, 300)
        print(SALIDA)


if __name__ == "__main__":
    main()
