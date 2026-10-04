"""Genera la presentación .pptx del tema 6 a partir de contenido.SLIDES."""
import math
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from contenido import SLIDES, M

# ------------------------------------------------------------------ estilo
INK = RGBColor(0x14, 0x2B, 0x33)      # texto
PRIM = RGBColor(0x0F, 0x4C, 0x5C)     # verde petróleo (dominante)
ACC = RGBColor(0xE3, 0x64, 0x14)      # naranja (acento)
TINT = RGBColor(0xE8, 0xF1, 0xF3)     # fondo suave
HL = RGBColor(0xFD, 0xE8, 0xD7)       # resaltado de filas
MUTED = RGBColor(0x5B, 0x6B, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ALT = RGBColor(0xF4, 0xF8, 0xF9)
HEAD = "Cambria"
BODY = "Calibri"
MATH = "Cambria Math"
MONO = "Courier New"

SW, SH = 13.333, 7.5
MX = 0.6
TOP = 1.55
BOTTOM = 6.9
GAP = 0.22

TOKEN = re.compile(r"\*\*|\^\{|_\{|\}")


def parse(text):
    """Devuelve lista de (texto, negrita, modo) con modo en '', 'sup', 'sub'."""
    out, bold, mode, pos = [], False, "", 0
    for m in TOKEN.finditer(text):
        tok = m.group()
        if tok == "}" and not mode:
            continue
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, mode))
        pos = m.end()
        if tok == "**":
            bold = not bold
        elif tok == "^{":
            mode = "sup"
        elif tok == "_{":
            mode = "sub"
        else:
            mode = ""
    if pos < len(text):
        out.append((text[pos:], bold, mode))
    return out


def plain(text):
    return "".join(t for t, _, _ in parse(text))


def add_runs(par, text, size, color=INK, font=BODY, bold=False):
    for t, b, mode in parse(text):
        r = par.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.name = font
        r.font.bold = bold or b
        r.font.color.rgb = color
        if mode:
            r.font._element.set("baseline", "30000" if mode == "sup" else "-25000")


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, margin=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tb, tf


def n_lines(text, w, size, factor=0.5):
    cpl = max(1, int(w * 72 / (size * factor)))
    return max(1, math.ceil(len(plain(text)) / cpl))


def lh(size):
    return size * 1.22 / 72


def card_cols(b):
    n = len(b["items"])
    if b.get("cols"):
        return b["cols"]
    return n if n <= 3 else (5 if n in (5, 10) else (4 if n == 8 else 3))


# ------------------------------------------------------------------ alturas
def height(b, w):
    t = b["t"]
    if b.get("h"):
        return b["h"]
    if t in ("bul", "num"):
        ind = 0.35
        return sum(n_lines(i, w - ind, b["size"], 0.47) * lh(b["size"]) + 0.09 for i in b["items"])
    if t == "txt":
        return n_lines(b["text"], w, b["size"]) * lh(b["size"]) + 0.04
    if t == "note":
        return n_lines(b["text"], w - 0.5, b["size"]) * lh(b["size"]) + 0.3
    if t == "form":
        return len(b["lines"]) * lh(b["size"]) * 1.12 + 0.3
    if t == "tab":
        return sum(row_heights(b, w))
    if t == "path":
        return 0.55
    if t == "mono":
        return (b["text"].count("\n") + 1) * b["size"] * 1.18 / 72 + 0.3
    if t == "cards":
        cols = card_cols(b)
        n = len(b["items"])
        rows = math.ceil(n / cols)
        cw = (w - (cols - 1) * 0.25) / cols
        ch = max(lh(b["size"] + 2) * n_lines(h_, cw - 0.4, b["size"] + 2, 0.55) + 0.12 +
                 n_lines(t_, cw - 0.4, b["size"]) * lh(b["size"]) + (0.45 if cols == 1 else 0.62) for h_, t_ in b["items"])
        return rows * ch + (rows - 1) * 0.25
    raise ValueError(t)


def col_widths(b, w):
    ncol = len(b["head"])
    if b["colw"]:
        return [w * c for c in b["colw"]]
    lens = [max([len(plain(b["head"][j])) * 1.3] + [len(plain(r[j])) for r in b["rows"]]) + 2.5 for j in range(ncol)]
    tot = sum(lens)
    return [w * l / tot for l in lens]


def row_heights(b, w):
    cw = col_widths(b, w)
    out = []
    for r in [b["head"]] + b["rows"]:
        nl = max(n_lines(c, cw[j] - 0.16, b["size"], 0.55 if r is b["head"] else 0.48) for j, c in enumerate(r))
        out.append(nl * lh(b["size"]) + b.get("pad", 0.13))
    return out


# ------------------------------------------------------------------ bloques
def draw_bul(slide, b, x, y, w, h, numbered=False):
    _, tf = textbox(slide, x, y, w, h)
    first = True
    for k, item in enumerate(b["items"]):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.35))))
        pPr.set("indent", str(-int(Inches(0.35))))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for el in pPr.findall(tag, pPr.nsmap):
                pPr.remove(el)
        from lxml import etree
        ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        clr = etree.SubElement(pPr, f"{{{ns}}}buClr")
        srgb = etree.SubElement(clr, f"{{{ns}}}srgbClr")
        srgb.set("val", "E36414")
        if numbered:
            fnt = etree.SubElement(pPr, f"{{{ns}}}buFont")
            fnt.set("typeface", "+mj-lt")
            au = etree.SubElement(pPr, f"{{{ns}}}buAutoNum")
            au.set("type", "arabicPeriod")
        else:
            fnt = etree.SubElement(pPr, f"{{{ns}}}buFont")
            fnt.set("typeface", "Arial")
            ch = etree.SubElement(pPr, f"{{{ns}}}buChar")
            ch.set("char", "•")
        p.space_after = Pt(6)
        if numbered and item.startswith("**"):
            add_runs(p, "\u200b", b["size"])
        add_runs(p, item, b["size"])


def draw_txt(slide, b, x, y, w, h):
    _, tf = textbox(slide, x, y, w, h)
    add_runs(tf.paragraphs[0], b["text"], b["size"], color=INK if b["size"] >= 15 else MUTED)


def draw_note(slide, b, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.12
    shp.fill.solid()
    shp.fill.fore_color.rgb = HL
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.22)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    add_runs(tf.paragraphs[0], b["text"], b["size"])


def draw_form(slide, b, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.1
    shp.fill.solid()
    shp.fill.fore_color.rgb = TINT
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    for k, line in enumerate(b["lines"]):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, line, b["size"], color=PRIM, font=MATH)


def draw_tab(slide, b, x, y, w, h):
    rows = [b["head"]] + b["rows"]
    nrow, ncol = len(rows), len(b["head"])
    cw = col_widths(b, w)
    rh = row_heights(b, w)
    gf = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(sum(rh)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    for j in range(ncol):
        tbl.columns[j].width = Inches(cw[j])
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(rh[i])
        for j, c in enumerate(r):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            p = cell.text_frame.paragraphs[0]
            sign = c in ("+", M)
            col_long = any(len(plain(rr[j])) > 30 for rr in b["rows"])
            p.alignment = PP_ALIGN.LEFT if (b["first_left"] and (j == 0 or col_long)) else PP_ALIGN.CENTER
            if i == 0:
                cell.fill.fore_color.rgb = PRIM
                add_runs(p, c, b["size"], color=WHITE, bold=True)
            else:
                cell.fill.fore_color.rgb = HL if (i - 1) in b["hl"] else (ALT if i % 2 == 0 else WHITE)
                if sign:
                    add_runs(p, c, b["size"] + 1, color=PRIM if c == "+" else ACC, bold=True)
                else:
                    add_runs(p, c, b["size"])


def draw_path(slide, b, x, y, w, h):
    cx = x
    for k, step in enumerate(b["steps"]):
        sw = len(step) * 0.125 + 0.45
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(y), Inches(sw), Inches(0.5))
        shp.adjustments[0] = 0.3
        shp.fill.solid()
        last = k == len(b["steps"]) - 1
        shp.fill.fore_color.rgb = ACC if last else PRIM
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], step, 16, color=WHITE, bold=True)
        cx += sw
        if not last:
            _, tf2 = textbox(slide, cx, y, 0.4, 0.5, anchor=MSO_ANCHOR.MIDDLE)
            tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
            add_runs(tf2.paragraphs[0], "›", 22, color=MUTED, bold=True)
            cx += 0.4


def draw_mono(slide, b, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = ALT
    shp.line.color.rgb = RGBColor(0xC9, 0xD6, 0xDA)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18)
    for k, line in enumerate(b["text"].split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = line if line else " "
        r.font.size = Pt(b["size"])
        r.font.name = MONO
        r.font.color.rgb = INK
        heading = line and not line.startswith(" ") and line.split()[0] in ("Estructura", "Análisis", "Resumen")
        r.font.bold = bool(heading)
        if heading:
            r.font.color.rgb = PRIM


def draw_cards(slide, b, x, y, w, h):
    n = len(b["items"])
    cols = card_cols(b)
    rows = math.ceil(n / cols)
    cw = (w - (cols - 1) * 0.25) / cols
    ch = (h - (rows - 1) * 0.25) / rows
    for k, (hd, tx) in enumerate(b["items"]):
        cx = x + (k % cols) * (cw + 0.25)
        cy = y + (k // cols) * (ch + 0.25)
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy), Inches(cw), Inches(ch))
        shp.adjustments[0] = 0.08
        shp.fill.solid()
        shp.fill.fore_color.rgb = TINT
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        add_runs(p, hd, b["size"] + 2, color=PRIM, font=HEAD, bold=True)
        p.space_after = Pt(6)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        add_runs(p2, tx, b["size"])


def style_chart(chart, b):
    chart.font.size = Pt(12)
    chart.font.name = BODY
    chart.font.color.rgb = INK
    if b.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = b["title"]
        r = chart.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = PRIM
    else:
        chart.has_title = False


def axis_title(axis, text):
    axis.has_title = True
    axis.axis_title.text_frame.text = text
    r = axis.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(12)
    r.font.bold = False
    r.font.name = BODY
    r.font.color.rgb = MUTED


def grid(axis, on):
    axis.has_major_gridlines = on
    if on:
        axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE5, 0xE8)
        axis.major_gridlines.format.line.width = Pt(0.5)
    axis.format.line.color.rgb = RGBColor(0x9A, 0xA8, 0xAE)


def draw_chart(slide, b, x, y, w, h):
    kind = b["kind"]
    if kind == "scatter":
        cd = XyChartData()
        s = cd.add_series("Datos")
        for px, py in b["points"]:
            s.add_data_point(px, py)
        if b.get("line") == "fit":
            xs = [p[0] for p in b["points"]]
            ys = [p[1] for p in b["points"]]
            n = len(xs)
            mx, my = sum(xs) / n, sum(ys) / n
            sl = sum((a - mx) * (c - my) for a, c in zip(xs, ys)) / sum((a - mx) ** 2 for a in xs)
            s2 = cd.add_series("Recta")
            for px in (min(xs), max(xs)):
                s2.add_data_point(px, my + sl * (px - mx))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        style_chart(chart, b)
        chart.has_legend = False
        ser = chart.plots[0].series[0]
        ser.format.line.fill.background()
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 9
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = PRIM
        ser.marker.format.line.color.rgb = WHITE
        for idx, lab in b["labels"].items():
            dl = ser.points[idx].data_label
            dl.has_text_frame = True
            dl.text_frame.text = lab
            dl.position = XL_LABEL_POSITION.RIGHT
            rr = dl.text_frame.paragraphs[0].runs[0]
            rr.font.size = Pt(13)
            rr.font.bold = True
            rr.font.color.rgb = ACC
        if b.get("line") == "fit":
            s2 = chart.plots[0].series[1]
            s2.marker.style = XL_MARKER_STYLE.NONE
            s2.format.line.color.rgb = ACC
            s2.format.line.width = Pt(1.5)
            s2.smooth = False
        axis_title(chart.category_axis, b["xlab"])
        axis_title(chart.value_axis, b["ylab"])
        grid(chart.value_axis, True)
        grid(chart.category_axis, False)
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        chart.value_axis.tick_labels.number_format = "0.0"
        chart.value_axis.tick_labels.number_format_is_linked = False
        xs = [p[0] for p in b["points"]]
        if set(xs) <= {-1, 1}:
            chart.category_axis.minimum_scale = -2
            chart.category_axis.maximum_scale = 2
            chart.category_axis.major_unit = 1
    elif kind == "bar":
        cd = CategoryChartData()
        cd.categories = [plain(c) for c in b["cats"]]
        cd.add_series("Valor", b["vals"])
        vertical = all(c.startswith("k = ") for c in b["cats"])
        ctype = XL_CHART_TYPE.COLUMN_CLUSTERED if vertical else XL_CHART_TYPE.BAR_CLUSTERED
        if not vertical:
            cd = CategoryChartData()
            cd.categories = [plain(c) for c in reversed(b["cats"])]
            cd.add_series("Valor", list(reversed(b["vals"])))
        gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        style_chart(chart, b)
        chart.has_legend = False
        plot = chart.plots[0]
        plot.gap_width = 55
        ser = plot.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = PRIM
        ser.invert_if_negative = False
        plot.has_data_labels = True
        plot.data_labels.number_format = b["fmt"]
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(12)
        plot.data_labels.font.color.rgb = INK
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        axis_title(chart.value_axis, b["xlab"])
        grid(chart.value_axis, True)
        grid(chart.category_axis, False)
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        chart.value_axis.tick_labels.number_format = "0"
        chart.value_axis.tick_labels.number_format_is_linked = False
    elif kind == "line":
        cd = CategoryChartData()
        cd.categories = b["cats"]
        for name, vals in b["series"]:
            cd.add_series(name, vals)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        style_chart(chart, b)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        for ser, col in zip(chart.plots[0].series, (PRIM, ACC)):
            ser.format.line.color.rgb = col
            ser.format.line.width = Pt(2.5)
            ser.smooth = False
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 9
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = col
            ser.marker.format.line.color.rgb = WHITE
        axis_title(chart.value_axis, b["ylab"])
        grid(chart.value_axis, True)
        grid(chart.category_axis, False)


CUBE = {"(1)": (0, 0, 0), "a": (1, 0, 0), "b": (0, 1, 0), "ab": (1, 1, 0),
        "c": (0, 0, 1), "ac": (1, 0, 1), "bc": (0, 1, 1), "abc": (1, 1, 1)}


def cube_xy(a, bb, c, x, y, w, h):
    """A horizontal, C vertical, B en profundidad."""
    side = min(w, h - 0.5) * 0.52
    dx, dy = side * 0.52, side * 0.38
    ox = x + (w - side - dx) / 2
    oy = y + dy + 0.25
    return ox + a * side + bb * dx, oy + (1 - c) * side - bb * dy


def draw_cube(slide, b, x, y, w, h):
    cap = 0.4 if b.get("caption") else 0
    hh = h - cap
    pts = {k: cube_xy(*v, x, y, w, hh) for k, v in CUBE.items()}
    names = list(CUBE)
    for i, p in enumerate(names):
        for q in names[i + 1:]:
            if sum(abs(u - v) for u, v in zip(CUBE[p], CUBE[q])) == 1:
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(pts[p][0]), Inches(pts[p][1]),
                                                Inches(pts[q][0]), Inches(pts[q][1]))
                ln.line.color.rgb = RGBColor(0x9A, 0xA8, 0xAE)
                ln.line.width = Pt(1.25)
    rad = 0.17
    for k, (px, py) in pts.items():
        sel = k in b["sel"]
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px - rad), Inches(py - rad), Inches(2 * rad), Inches(2 * rad))
        c.fill.solid()
        c.fill.fore_color.rgb = ACC if sel else WHITE
        c.line.color.rgb = ACC if sel else RGBColor(0x9A, 0xA8, 0xAE)
        c.line.width = Pt(1.5)
        c.shadow.inherit = False
        right = CUBE[k][0] == 1
        _, tf = textbox(slide, px + (0.22 if right else -0.92), py - 0.36, 0.7, 0.32)
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT
        add_runs(tf.paragraphs[0], k, 15, color=ACC if sel else MUTED, bold=sel, font=MATH)
    if cap:
        _, tf = textbox(slide, x, y + hh, w, cap)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], b["caption"], 13, color=MUTED)


DRAW = {"bul": draw_bul, "num": lambda s, b, x, y, w, h: draw_bul(s, b, x, y, w, h, numbered=True),
        "txt": draw_txt, "note": draw_note, "form": draw_form, "tab": draw_tab, "path": draw_path,
        "mono": draw_mono, "cards": draw_cards, "chart": draw_chart, "cube": draw_cube}

WARN = []


def column(slide, blocks, x, w, title):
    hs = [height(b, w) for b in blocks]
    total = sum(hs) + GAP * (len(blocks) - 1)
    avail = BOTTOM - TOP
    if total > avail + 0.01:
        WARN.append(f"DESBORDE {total - avail:+.2f} in: {plain(title)}")
    y = TOP
    for b, h in zip(blocks, hs):
        DRAW[b["t"]](slide, b, x, y, w, h)
        y += h + GAP


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def footer(slide, sec, n):
    _, tf = textbox(slide, MX, 7.05, 8, 0.3)
    add_runs(tf.paragraphs[0], f"Diseños factoriales fraccionados · {sec}" if sec else "Diseños factoriales fraccionados",
             10, color=MUTED)
    _, tf = textbox(slide, SW - MX - 1, 7.05, 1, 0.3)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_runs(tf.paragraphs[0], str(n), 10, color=MUTED)


def set_title(slide, text, color=PRIM, size=30, y=0.42, h=1.0, anchor=MSO_ANCHOR.MIDDLE):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(MX), Inches(y), Inches(SW - 2 * MX), Inches(h)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, text, size, color=color, font=HEAD, bold=True)


def build(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    lay = prs.slide_layouts[5]  # solo título
    for n, s in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(lay)
        kind = s["kind"]
        if kind in ("title", "closing"):
            bg(slide, PRIM)
            set_title(slide, s["title"], color=WHITE, size=46, y=2.3, h=1.5, anchor=MSO_ANCHOR.BOTTOM)
            _, tf = textbox(slide, MX, 4.0, SW - 2 * MX, 0.6)
            add_runs(tf.paragraphs[0], s["sub"], 22, color=RGBColor(0xCF, 0xE3, 0xE8))
            if s.get("meta"):
                _, tf = textbox(slide, MX, 5.3, SW - 2 * MX, 1.0)
                for k, line in enumerate(s["meta"].split("\n")):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    add_runs(p, line, 16, color=WHITE)
            # motivo: signos + y −
            for k, (ch, col) in enumerate([("+", ACC), (M, WHITE), ("+", WHITE), (M, ACC)]):
                c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(SW - 3.6 + (k % 2) * 1.3),
                                           Inches(0.55 + (k // 2) * 1.2), Inches(1.0), Inches(1.0))
                c.fill.solid()
                c.fill.fore_color.rgb = col
                c.line.fill.background()
                c.shadow.inherit = False
                tf = c.text_frame
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                add_runs(tf.paragraphs[0], ch, 36, color=PRIM if col == WHITE else WHITE, bold=True)
        elif kind == "section":
            bg(slide, PRIM)
            c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(MX), Inches(2.2), Inches(1.3), Inches(1.3))
            c.fill.solid()
            c.fill.fore_color.rgb = ACC
            c.line.fill.background()
            c.shadow.inherit = False
            tf = c.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            add_runs(tf.paragraphs[0], str(s["n"]), 40, color=WHITE, font=HEAD, bold=True)
            set_title(slide, s["title"], color=WHITE, size=44, y=3.7, h=1.0, anchor=MSO_ANCHOR.TOP)
            _, tf = textbox(slide, MX, 4.75, SW - 2 * MX, 0.6)
            add_runs(tf.paragraphs[0], s["sub"], 20, color=RGBColor(0xCF, 0xE3, 0xE8))
        else:
            bg(slide, WHITE)
            set_title(slide, s["title"])
            W = SW - 2 * MX
            if s.get("body"):
                column(slide, s["body"], MX, W, s["title"])
            else:
                g = 0.45
                lw = (W - g) * s["ratio"]
                column(slide, s["left"], MX, lw, s["title"])
                column(slide, s["right"], MX + lw + g, W - g - lw, s["title"])
            footer(slide, s["sec"], n)
        if s.get("explica"):
            slide.notes_slide.notes_text_frame.text = plain(s["explica"])
    prs.save(out)
    for w_ in WARN:
        print(w_)
    print("diapositivas:", len(SLIDES))


if __name__ == "__main__":
    build(sys.argv[1])
